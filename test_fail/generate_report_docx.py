import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" />')
    tcPr.append(shd)

def create_report():
    doc = Document()
    
    # 1. Margins setup (Top/Bottom 2cm, Left 3cm, Right 1.5cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.59)
        
    # Default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BÁO CÁO ĐÁNH GIÁ TIẾN ĐỘ VÀ THỰC TRẠNG DỰ ÁN (GIAI ĐOẠN 1)\n(Khai thác chi tiết từ các phiên thảo luận chuyên sâu & Lịch sử dự án)")
    r_title.font.bold = True
    r_title.font.size = Pt(15)
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Dự án: Hệ thống AI Sinh Giáo Án Chuẩn Công văn 5512 & Slide Bài Giảng Tự Động (LessonAI-5512)\n")
    r_sub.font.italic = True
    r_sub.font.size = Pt(12)
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.add_run("―" * 50)
    
    # SECTION I: ĐÁNH GIÁ MỨC ĐỘ HOÀN THÀNH DỰ ÁN HIỆN TẠI
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("I. ĐÁNH GIÁ MỨC ĐỘ HOÀN THÀNH DỰ ÁN HIỆN TẠI")
    r_h1.font.bold = True
    r_h1.font.size = Pt(13)
    r_h1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.add_run(
        "Nhìn nhận một cách khách quan và thẳng thắn từ thực tế mã nguồn và các thảo luận trong lịch sử dự án, đồ án hiện tại mới chỉ đạt mức "
    )
    r_b1 = p1.add_run("Proof of Concept (POC) / Mức độ 'Chạy được Script & API cơ bản'")
    r_b1.font.bold = True
    p1.add_run(
        ", tương đương khoảng 30% - 35% khối lượng của một hệ thống sản phẩm hoàn chỉnh. "
        "Dự án mới chỉ chứng minh được tính khả thi về mặt thuật toán: Chuyển dữ liệu từ AI -> JSON -> File Word & PowerPoint. "
        "Dự án tuyệt đối CHƯA THỂ xem là một sản phẩm hoàn thiện sẵn sàng cho giáo viên sử dụng (Production-ready).\n"
    )
    
    # Summary Table of Completion Level
    table_comp = doc.add_table(rows=6, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_comp.rows[0].cells
    hdr[0].text = "Thành phần Hệ thống"
    hdr[1].text = "Trạng thái Hiện tại"
    hdr[2].text = "Mức độ Hoàn thành"
    for c in hdr:
        set_cell_background(c, "003366")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    rows_data = [
        ("AI Generation Engine (Gemini API)", "Đã gọi được API Gemini 2.5 Flash xuất JSON theo Pydantic Schema", "Đạt ~60% (Chưa streaming, phụ thuộc API 3rd party)"),
        ("Bộ Kiểm duyệt 5512 (Validator)", "Đã có script check từ khóa và đếm độ dài bước thô sơ", "Đạt ~30% (Chưa đánh giá chiều sâu sư phạm)"),
        ("Module Xuất File Word (.docx)", "Render được công thức LaTeX sang OMML XML native trong Word", "Đạt ~65% (Tự phát triển hàm convert LaTeX -> OMML)"),
        ("Module Xuất Slide PPTX", "Tự động xuất Slide PPTX qua công cụ Marp CLI", "Đạt ~40% (Slide còn thô sơ, chưa tùy biến visual)"),
        ("Cơ sở dữ liệu (Database) & UI", "Hoàn toàn chưa triển khai (Mới chạy mô hình Prompt / REST API)", "Đạt 0% (Chưa có CSDL & Giao diện người dùng)")
    ]
    for idx, row in enumerate(rows_data, start=1):
        cells = table_comp.rows[idx].cells
        cells[0].text = row[0]
        cells[1].text = row[1]
        cells[2].text = row[2]
        if idx % 2 == 1:
            for c in cells:
                set_cell_background(c, "F2F4F7")
                
    doc.add_paragraph()
    
    # SECTION II: PHÂN TÍCH TỪ CÁC THẢO LUẬN CHUYÊN SÂU TRONG LỊCH SỬ DỰ ÁN
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("II. PHÂN TÍCH VẤN ĐỀ TỪ LỊCH SỬ THẢO LUẬN CHUYÊN SÂU (THREAD 'QUESTION' & THẮC MẮC KỸ THUẬT)")
    r_h2.font.bold = True
    r_h2.font.size = Pt(13)
    r_h2.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p2 = doc.add_paragraph()
    p2.add_run("Trong quá trình phát triển đồ án, các chủ đề cốt lõi đã được đặt ra và phân tích trong lịch sử trao đổi bao gồm:")
    
    topics = [
        ("1. Tranh luận: Dự án AI chỉ cần 'Prompt + Result' hay bắt buộc cần Cơ sở dữ liệu (Database)?",
         "• Thắc mắc đã đặt ra: 'Tại sao lại cần CSDL? Nếu áp dụng AI không phải chỉ cần đơn giản là Prompt và kết quả nhận lại từ AI sao? Phát triển dạng Chatbot có cần CSDL không?'\n"
         "• Thực tế đồ án hiện tại: Đồ án hiện đang chạy thuần theo hướng PURE PROMPT & API (không có CSDL). Người dùng gửi yêu cầu -> Backend gọi Gemini -> Xuất file.\n"
         "• Đánh giá hạn chế: Chính việc CHƯA CÓ CSDL làm cho dự án chỉ dừng lại ở mức POC 'chạy được'. Hệ thống không lưu giữ được lịch sử bài dạy của giáo viên, không phân quyền người dùng, không quản lý được bộ Template slide và không thể theo dõi phiên làm việc dài hạn."),
        
        ("2. Vấn đề Chuyển đổi Công thức Toán LaTeX sang Microsoft Word (LaTeX to OMML Conversion)",
         "• Thắc mắc đã đặt ra: 'Làm sao khi kết quả trả về là file Word nhưng trong đó các công thức toán học được đưa về dạng chuẩn Word chứ không bị vỡ ảnh hay hiển thị dạng chuỗi LaTeX thô?'\n"
         "• Thực tế đã giải quyết: Đã xây dựng hàm convert_latex_to_omml() trong docx_exporter.py kết hợp latex2mathml + mathml2omml + parse_xml để tự động biến đổi chuỗi \\(...\\), $...$ thành OMML Native Equation trong Word. Đây là một điểm sáng xử lý kỹ thuật tốt trong đồ án."),
        
        ("3. Vấn đề Slide Bài Giảng thô sơ & Hạn chế của Marp CLI",
         "• Thắc mắc đã đặt ra: 'Tại sao khi để AI sinh ra slide, slide sinh ra như không thể chỉnh sửa linh hoạt, bị gò bó? Marp CLI chạy trên môi trường Node.js có ổn định khi triển khai thực tế không?'\n"
         "• Thực tế hiện tại: Đã tích hợp Marp CLI biên dịch Markdown sang PPTX. Tuy nhiên, Slide sinh ra còn rất thô, mang phong cách ghi chú văn bản, chưa có hình ảnh minh họa thật (mới chỉ có text image_prompt) và chưa can thiệp sâu vào layout PowerPoint native qua python-pptx."),
        
        ("4. Vấn đề Đánh giá Logic Quy chuẩn 5512 & Tính 'Kỷ luật hóa' AI",
         "• Thắc mắc đã đặt ra: 'Xử lý AI và logic 5512 là những gì? Thể hiện qua các dòng code như thế nào? Liệu code của tôi đã giải quyết triệt để hay chưa?'\n"
         "• Thực tế hiện tại: Đã áp dụng Structured Output (Pydantic Schema) ép AI ra JSON đúng khung 4 hoạt động & 4 bước tổ chức. Nhưng bộ kiểm duyệt validation.py mới chỉ là rule-based thô (check từ khóa & độ dài chuỗi > 10 ký tự), chưa đánh giá được chiều sâu sư phạm hay tính chính xác kiến thức.")
    ]
    
    for title, body in topics:
        p_top = doc.add_paragraph()
        r_tt = p_top.add_run(title)
        r_tt.font.bold = True
        r_tt.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
        
        p_bd = doc.add_paragraph()
        p_bd.paragraph_format.left_indent = Inches(0.2)
        p_bd.paragraph_format.line_spacing = 1.15
        p_bd.add_run(body)
        doc.add_paragraph()
        
    # SECTION III: CHI TIẾT CÁC CÔNG VIỆC ĐÃ XỬ LÝ VÀ ĐÁNH GIÁ MỨC ĐỘ (CƠ BẢN/NÂNG CAO)
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("III. CHI TIẾT CÁC CÔNG VIỆC ĐÃ XỬ LÝ VÀ ĐÁNH GIÁ MỨC ĐỘ")
    r_h3.font.bold = True
    r_h3.font.size = Pt(13)
    r_h3.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    tasks = [
        ("1. Thiết kế Pydantic Data Schemas (schemas.py, slide_schema.py)", "CƠ BẢN", 
         "Định nghĩa cấu trúc JSON chuẩn cho Giáo án 5512 và Slide bài giảng. Đây là bước thiết kế khung dữ liệu tiêu chuẩn."),
        ("2. Tích hợp Gemini API & Structured Output (ai_service.py)", "CƠ BẢN - TRUNG BÌNH",
         "Sử dụng Google GenAI SDK (gemini-2.5-flash) với response_schema ép AI trả về dữ liệu chuẩn JSON, hạn chế vỡ cấu trúc."),
        ("3. Xây dựng Rule-based Validator 5512 (validation.py)", "CƠ BẢN",
         "Viết hàm kiểm tra quy chuẩn bằng từ khóa và độ dài chuỗi văn bản. Chưa có AI kiểm duyệt nội dung sâu."),
        ("4. Cơ chế Sửa lỗi tự động Auto-Retry Loop (ai_service.py)", "TRUNG BÌNH",
         "Tự động ghép phản hồi lỗi của Validator vào Prompt gửi lại cho Gemini sửa (tối đa 2 lần)."),
        ("5. Render Công thức Toán Native trong Word (docx_exporter.py)", "KHÁ - NÂNG CAO",
         "Tự phát triển module bóc tách LaTeX và chuyển sang OMML XML Native hiển thị sắc nét trong Word. Lọc sạch ký tự rác XML."),
        ("6. Chuyển đổi Slide Marp & Xuất PowerPoint (marp_converter.py, slide_exporter.py)", "CƠ BẢN",
         "Chuyển JSON Slide sang Marp Markdown và gọi Marp CLI xuất PPTX."),
        ("7. Đóng gói REST API Server (main.py)", "CƠ BẢN",
         "Tạo 4 endpoint FastAPI cơ bản nhận bài dạy và trả về file byte stream (.docx, .pptx). Chưa có Auth, chưa có Async Streaming.")
    ]
    
    for title, level, desc in tasks:
        p_t = doc.add_paragraph()
        r_t1 = p_t.add_run(f"• {title} - ")
        r_t1.font.bold = True
        r_t2 = p_t.add_run(f"[{level}]\n")
        r_t2.font.bold = True
        r_t2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) if "CƠ BẢN" in level else RGBColor(0x00, 0x80, 0x00)
        
        p_d = doc.add_paragraph()
        p_d.paragraph_format.left_indent = Inches(0.25)
        p_d.add_run(desc)
        
    doc.add_paragraph()
    
    # SECTION IV: CÁC KỸ THUẬT ĐÃ SỬ DỤNG
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("IV. TỔNG HỢP CÁC KỸ THUẬT VÀ CÔNG NGHỆ ĐÃ SỬ DỤNG")
    r_h4.font.bold = True
    r_h4.font.size = Pt(13)
    r_h4.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    techs = [
        ("Backend Framework", "Python 3.10+, FastAPI, Uvicorn Server, CORS Middleware."),
        ("AI / LLM Orchestration", "Google GenAI SDK (google-genai), Gemini 2.5 Flash Model, Structured Output (response_schema)."),
        ("Validation & Schema", "Pydantic v2 (BaseModel, Field, model_dump), Custom Rule-Based Regex Matching."),
        ("Word Math Engine", "python-docx, latex2mathml, mathml2omml, parse_xml (LaTeX -> OMML Native Equations)."),
        ("Slide Generator Engine", "Marp CLI (@marp-team/marp-cli via Node.js subprocess), Marp Markdown Syntax, CSS Grid Layout."),
        ("System Utility", "python-dotenv, io.BytesIO (Memory streaming cho file download).")
    ]
    
    table_tech = doc.add_table(rows=len(techs)+1, cols=2)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = table_tech.rows[0].cells
    t_hdr[0].text = "Hạng mục Công nghệ"
    t_hdr[1].text = "Chi tiết Kỹ thuật đã áp dụng"
    for c in t_hdr:
        set_cell_background(c, "003366")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    for idx, (cat, spec) in enumerate(techs, start=1):
        cells = table_tech.rows[idx].cells
        cells[0].text = cat
        cells[1].text = spec
        if idx % 2 == 1:
            for c in cells:
                set_cell_background(c, "F2F4F7")
                
    doc.add_paragraph()
    
    # SECTION V: ƯU ĐIỂM VÀ NHƯỢC ĐIỂM Ở GIAI ĐOẠN HIỆN TẠI
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("V. ĐÁNH GIÁ ƯU ĐIỂM VÀ NHƯỢC ĐIỂM CỦA GIAI ĐOẠN HIỆN TẠI")
    r_h5.font.bold = True
    r_h5.font.size = Pt(13)
    r_h5.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_pro = doc.add_paragraph()
    r_pro = p_pro.add_run("1. Ưu điểm (Những kết quả tích cực đã đạt được):")
    r_pro.font.bold = True
    
    pros = [
        "Đã chạy thông suốt luồng dữ liệu tự động từ Yêu cầu người dùng -> LLM JSON -> Validation -> File .docx & .pptx.",
        "Giải quyết triệt để bài toán render công thức Toán LaTeX thành công thức OMML Native Equation trong Word.",
        "Áp dụng thành công Structured Output giúp kết quả AI không bị lệch khung cấu trúc JSON 5512.",
        "Kiến trúc Backend được chia mô-đun rõ ràng, dễ bảo trì và mở rộng."
    ]
    for pr in pros:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.add_run(f"✔  {pr}")
        
    p_con = doc.add_paragraph()
    r_con = p_con.add_run("\n2. Nhược điểm & Hạn chế lớn tồn tại:")
    r_con.font.bold = True
    
    cons = [
        "Hoàn toàn CHƯA CÓ GIAO DIỆN NGƯỜI DÙNG (No UI): Người dùng phải gọi qua API/Terminal, giáo viên chưa thể sử dụng.",
        "Chất lượng Slide PowerPoint từ Marp rất thô sơ: Thiếu tính thẩm mỹ sư phạm, không có bố cục màu sắc linh hoạt, chưa chèn được hình ảnh thật.",
        "Thiếu Cơ sở dữ liệu (Database): Không lưu trữ được lịch sử giáo án, không quản lý tài khoản hay phiên làm việc.",
        "Bộ kiểm duyệt 5512 quá máy móc: Chỉ kiểm tra từ khóa bề nổi, chưa đánh giá được chất lượng nội dung sư phạm thực tế.",
        "Thời gian phản hồi chậm (10-20s): Do gọi LLM và retry đồng bộ, chưa có cơ chế Async Streaming.",
        "Phụ thuộc 100% vào Gemini API: Chưa có phương án dự phòng khi API bị nghẽn hay hết quota."
    ]
    for cn in cons:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.add_run(f"✖  {cn}")
        
    doc.add_paragraph()
    
    # SECTION VI: CÁC VẤN ĐỀ ĐÃ GIẢI QUYẾT VÀ TỒN ĐỌC CHƯA GIẢI QUYẾT
    h6 = doc.add_paragraph()
    r_h6 = h6.add_run("VI. TỔNG HỢP VẤN ĐỀ ĐÃ GIẢI QUYẾT VÀ TỒN ĐỌC CHƯA GIẢI QUYẾT")
    r_h6.font.bold = True
    r_h6.font.size = Pt(13)
    r_h6.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    table_issues = doc.add_table(rows=6, cols=2)
    table_issues.alignment = WD_TABLE_ALIGNMENT.CENTER
    i_hdr = table_issues.rows[0].cells
    i_hdr[0].text = "ĐÃ GIẢI QUYẾT ĐƯỢC (Solved)"
    i_hdr[1].text = "TỒN ĐỌC CHƯA GIẢI QUYẾT (Unsolved / Pending)"
    for c in i_hdr:
        set_cell_background(c, "003366")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    issues_data = [
        ("Ép AI trả về đúng cấu trúc JSON Giáo án 5512 bằng Pydantic Schema.", "Chưa có Giao diện Web (UI) cho giáo viên nhập liệu và chỉnh sửa."),
        ("Render công thức toán LaTeX thành công thức OMML chuẩn trong file Word.", "Slide PPTX sinh qua Marp CLI còn quá thô, chưa dùng python-pptx native template."),
        ("Lọc sạch ký tự rác XML ẩn tránh hỏng file .docx khi xuất.", "Chưa chèn được hình ảnh thật vào Slide (mới chỉ có text image_prompt)."),
        ("Xây dựng bộ kiểm duyệt từ khóa và cấu trúc 4 hoạt động 5512 cơ bản.", "Validator chưa đánh giá được chất lượng nội dung sư phạm thực tế."),
        ("Tự động gửi feedback cho AI sửa lại khi Validator phát hiện sai cấu trúc.", "Thiếu Database lưu trữ tài khoản, lịch sử giáo án và quản lý phiên làm việc.")
    ]
    for idx, (sol, unsol) in enumerate(issues_data, start=1):
        cells = table_issues.rows[idx].cells
        cells[0].text = f"✔ {sol}"
        cells[1].text = f"✖ {unsol}"
        if idx % 2 == 1:
            for c in cells:
                set_cell_background(c, "F2F4F7")
                
    doc.add_paragraph()
    
    # SECTION VII: KẾT QUẢ ĐÃ VÀ ĐÀNG NHẬN ĐƯỢC THỰC CHẤT THỂ HIỆN ĐIỀU GÌ?
    h7 = doc.add_paragraph()
    r_h7 = h7.add_run("VII. KẾT QUẢ ĐÃ VÀ ĐÀNG NHẬN ĐƯỢC THỰC CHẤT THỂ HIỆN ĐIỀU GÌ?")
    r_h7.font.bold = True
    r_h7.font.size = Pt(13)
    r_h7.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_ans = doc.add_paragraph()
    p_ans.paragraph_format.line_spacing = 1.15
    p_ans.add_run(
        "Đánh giá một cách nghiêm túc, thẳng thắn và sát với thực tế dự án:\n\n"
    )
    
    points = [
        ("1. CHỈ THỂ HIỆN TÍNH KHẢ THI VỀ KỸ THUẬT (Proof of Concept - POC):", 
         "Kết quả chạy được hiện tại chứng minh rằng ý tưởng kết hợp AI + Validator + Document Exporters là KHẢ THI về mặt thuật toán. Luồng dữ liệu có thể chảy tự động từ AI ra file Word và PowerPoint mà không bị đứt gãy."),
        ("2. CHƯA THỂ HIỆN TÍNH SẴN SÀNG SỬ DỤNG (Not Production-Ready):", 
         "Dự án tuyệt đối CHƯA THỂ coi là hoàn thành. Thiếu UI, thiếu CSDL, Slide xấu và Validator thô sơ khiến dự án hiện tại mới chỉ là 'khung xương kỹ thuật thô' (Basic Backend Skeleton), chưa thể đưa cho bất kỳ giáo viên nào sử dụng thực tế."),
        ("3. ĐÓNG VAI TRÒ LÀM NỀN TẢNG ĐỂ PHÁT TRIỂN TIẾP:", 
         "Những gì làm được đến lúc này đóng vai trò làm móng: xác nhận mô hình dữ liệu đúng, thuật toán xuất file chạy được. Tất cả phần việc quyết định giá trị thực sự của đồ án (Giao diện Web, Thiết kế Slide đẹp, Đánh giá chất lượng Sư phạm sâu, CSDL) nằm ở các giai đoạn phía trước.")
    ]
    
    for title_p, desc_p in points:
        p_pt = doc.add_paragraph()
        r_p1 = p_pt.add_run(f"{title_p}\n")
        r_p1.font.bold = True
        r_p1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        
        p_pd = doc.add_paragraph()
        p_pd.paragraph_format.left_indent = Inches(0.25)
        p_pd.add_run(desc_p)
        doc.add_paragraph()
        
    output_path = "/Users/nguyennhan18/Documents/Lesson-Plan-Generator/report_GD1.docx"
    doc.save(output_path)
    print(f"✅ Enhanced report saved successfully to {output_path}")

if __name__ == "__main__":
    create_report()
