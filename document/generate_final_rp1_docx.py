import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" />')
    tcPr.append(shd)

def create_final_rp1():
    doc = Document()
    
    # Standard Margins (Top/Bottom 2cm, Left 3cm, Right 1.5cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.59)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # -------------------------------------------------------------
    # DOCUMENT TITLE
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BÁO CÁO TỔNG KẾT GIAI ĐOẠN 1 (FINAL REPORT - STAGE 1)\n")
    r_title.font.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Dự án: Hệ thống AI Sinh Giáo Án Chuẩn Công văn 5512 & Slide Bài Giảng Tự Động (LessonAI-5512)\n")
    r_sub.font.italic = True
    r_sub.font.size = Pt(12)
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Báo cáo tổng hợp: Đã thực hiện, Kỹ thuật áp dụng, Phân tích vấn đề từ Thread 'Question' & Hướng giải quyết\n")
    r_meta.font.bold = True
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.add_run("―" * 55)
    
    # -------------------------------------------------------------
    # SECTION I: TỔNG QUAN NHỮNG GÌ ĐÃ THỰC HIỆN VÀ XỬ LÝ ĐƯỢC
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("I. TỔNG QUAN CÁC CÔNG VIỆC ĐÃ THỰC HIỆN VÀ XỬ LÝ ĐƯỢC TRONG GIAI ĐOẠN 1")
    r_h1.font.bold = True
    r_h1.font.size = Pt(13)
    r_h1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.add_run(
        "Trong Giai đoạn 1, dự án đã xây dựng thành công bộ "
    )
    r_b1 = p1.add_run("Khung Backend thử nghiệm thô (Proof of Concept - POC)")
    r_b1.font.bold = True
    p1.add_run(
        " xử lý luồng dữ liệu tự động từ lúc tiếp nhận Yêu cầu bài dạy của giáo viên cho tới khi xuất ra hai định dạng tài liệu chuẩn: "
        "File Word (.docx) Kế hoạch bài dạy theo Công văn 5512 và File PowerPoint (.pptx) Slide trình chiếu bài giảng. "
        "Chi tiết các công việc đã thực hiện gồm:\n"
    )
    
    tasks_done = [
        ("1. Thiết kế Schema Dữ liệu Chuẩn hóa (schemas.py, slide_schema.py)",
         "Xây dựng bộ Pydantic Models định nghĩa cấu trúc dữ liệu nghiêm ngặt cho Kế hoạch bài dạy 5512 (bao gồm: Thông tin chung, Mục tiêu Kiến thức/Năng lực, 4 Hoạt động học tập, 4 Bước tổ chức thực hiện) và Slide Deck (Title Slide, Warm-up, Concept, Exercise, Summary)."),
        
        ("2. Động cơ Tích hợp AI LLM Gemini 2.5 Flash (ai_service.py)",
         "Sử dụng Google GenAI SDK (mô hình gemini-2.5-flash) kết hợp với kỹ thuật Structured Output (response_mime_type='application/json' & response_schema) để ép AI sinh dữ liệu đúng 100% định dạng JSON Schema, không bị lệch khung hay vỡ cấu trúc."),
        
        ("3. Bộ Kiểm duyệt Quy chuẩn 5512 & Auto-Retry Loop (validation.py & ai_service.py)",
         "Phát triển module validate_5512_lesson_plan() tự động quét dữ liệu đầu ra: kiểm tra số lượng hoạt động (=4), từ khóa chuẩn 5512 ('mở đầu/khởi động', 'hình thành kiến thức', 'luyện tập', 'vận dụng') và độ dài 4 bước thực hiện. Nếu phát hiện lỗi, hệ thống tự động nối feedback lỗi vào Prompt để yêu cầu Gemini sửa lại (tối đa 2 lần)."),
        
        ("4. Module Render Công thức Toán Native trong Microsoft Word (docx_exporter.py)",
         "Xây dựng thành công thuật toán bóc tách công thức toán LaTeX \\(...\\), $...$ và chuyển đổi 3 bước: LaTeX -> MathML -> OMML XML (Word Native Equation). Đồng thời viết bộ lọc clean_xml_string() triệt hạ 100% ký tự rác XML ẩn gây hỏng file Word."),
        
        ("5. Module Chuyển đổi Slide Marp & Xuất PowerPoint (marp_converter.py, slide_exporter.py)",
         "Phát triển hàm ánh xạ tự động từ Lesson Plan JSON sang Marp Markdown (hỗ trợ layout 2 cột, theme gaia/default/unwind, KaTeX math, comment kịch bản giảng dạy cho GV), sau đó gọi Marp CLI qua subprocess biên dịch sang file PPTX."),
        
        ("6. Đóng gói RESTful API Server bằng FastAPI (main.py)", "Tạo 4 REST endpoints (/api/v1/lesson-plan/generate, /api/v1/lesson-plan/export-docx, /api/v1/slide/generate, /api/v1/lesson-plan/export-slide) sẵn sàng nhận request và trả về byte stream file download.")
    ]
    
    for t_title, t_desc in tasks_done:
        p_t = doc.add_paragraph()
        r_tt = p_t.add_run(f"• {t_title}\n")
        r_tt.font.bold = True
        
        p_td = doc.add_paragraph()
        p_td.paragraph_format.left_indent = Inches(0.25)
        p_td.paragraph_format.line_spacing = 1.15
        p_td.add_run(t_desc)
        doc.add_paragraph()
        
    # -------------------------------------------------------------
    # SECTION II: CÁC KỸ THUẬT VÀ CÔNG NGHỆ ĐÃ SỬ DỤNG
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("II. BẢNG TỔNG HỢP CÁC KỸ THUẬT VÀ CÔNG NGHỆ ĐÃ SỬ DỤNG")
    r_h2.font.bold = True
    r_h2.font.size = Pt(13)
    r_h2.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    table_tech = doc.add_table(rows=7, cols=3)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = table_tech.rows[0].cells
    t_hdr[0].text = "Lĩnh vực Kỹ thuật"
    t_hdr[1].text = "Công nghệ / Thư viện Sử dụng"
    t_hdr[2].text = "Vai trò Kỹ thuật trong Dự án"
    for c in t_hdr:
        set_cell_background(c, "003366")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    techs = [
        ("Backend Framework", "Python 3.10+, FastAPI, Uvicorn ASGI Server, CORS Middleware", "Khởi tạo RESTful API Server bất đồng bộ, mở kết nối CORS cho Frontend Web."),
        ("AI / LLM Orchestration", "Google GenAI SDK (google-genai), Gemini 2.5 Flash, Structured Output", "Xử lý ngôn ngữ tự nhiên, hiểu ngữ cảnh sư phạm tiếng Việt, sinh JSON chuẩn schema."),
        ("Data Schema & Validation", "Pydantic v2 (BaseModel, Field), Custom Rule-Based Regex", "Định nghĩa cấu trúc dữ liệu chặt chẽ và kiểm tra quy chuẩn sư phạm 5512."),
        ("Mathematics Rendering", "latex2mathml, mathml2omml, lxml parse_xml, Regular Expression", "Biến đổi chuỗi LaTeX thành công thức OMML Native Equation hiển thị sắc nét trong MS Word."),
        ("Slide Generation Engine", "Marp CLI (@marp-team/marp-cli via Node.js/npx subprocess)", "Biên dịch nội dung Marp Markdown tự động sang file PowerPoint (.pptx)."),
        ("System Utility", "python-dotenv, io.BytesIO", "Quản lý biến môi trường API Key và xử lý stream bộ nhớ cho luồng tải file.")
    ]
    
    for idx, (cat, tech_item, role_desc) in enumerate(techs, start=1):
        cells = table_tech.rows[idx].cells
        cells[0].text = cat
        cells[1].text = tech_item
        cells[2].text = role_desc
        if idx % 2 == 1:
            for c in cells:
                set_cell_background(c, "F2F4F7")
                
    doc.add_paragraph()
    
    # -------------------------------------------------------------
    # SECTION III: PHÂN TÍCH VẤN ĐỀ TỪ THREAD 'QUESTION' & THẮC MẮC TRONG LỊCH SỬ CHAT
    # -------------------------------------------------------------
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("III. PHÂN TÍCH VẤN ĐỀ ĐANG MẮC PHẢI DỰA TRÊN THREAD 'QUESTION' & LỊCH SỬ CHAT")
    r_h3.font.bold = True
    r_h3.font.size = Pt(13)
    r_h3.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p3 = doc.add_paragraph()
    p3.paragraph_format.line_spacing = 1.15
    p3.add_run(
        "Khai thác từ các thắc mắc cốt lõi của tác giả trong thread "
    )
    r_q_b = p3.add_run("'question'")
    r_q_b.font.bold = True
    p3.add_run(
        " (và các thread kỹ thuật liên quan như a2663ef2, fa1dd4cc), báo cáo bóc tách 5 VẤN ĐỀ NGHẼN CỐT LÕI mà đồ án đang mắc phải trong Giai đoạn 1:\n\n"
    )
    
    issues_analyzed = [
        ("1. Vấn đề 'Nhầm lẫn Kiến trúc': Dự án chỉ cần 'Prompt + Result' hay phải có Cơ sở dữ liệu (Database)?",
         "• Thắc mắc đã đặt ra trong thread 'question': 'Tại sao lại cần CSDL? Nếu ta áp dụng AI để làm việc không phải chỉ cần đơn giản là Prompt và kết quả nhận lại từ AI sao? Nếu phát triển theo hướng Chatbot thì có cần CSDL không?'\n"
         "• Bóc tách nguyên nhân & Vấn đề đang mắc phải: Mã nguồn Giai đoạn 1 hiện tại ĐANG CHẠY THUẦN PURE PROMPT / API KHÔNG CSDL. Mỗi lần người dùng gửi request, Backend gọi Gemini rồi xuất file trực tiếp. Việc THIẾU CSDL làm cho hệ thống không lưu được lịch sử giáo án của giáo viên, không có quản lý tài khoản, không lưu vết phiên làm việc và không thể quản lý bộ mẫu Template Slide. Đây là nguyên nhân hàng đầu khiến dự án bị kẹt ở mức 'Script POC' chứ chưa thể thành sản phẩm phần mềm thực sự.\n"
         "• Hướng giải quyết đề xuất: Cần bổ sung ngay Cơ sở dữ liệu (PostgreSQL / MongoDB) kết hợp SQLAlchemy ORM ở Giai đoạn 2 để quản lý Users, Saved Lesson Plans, Slide Templates và Session History."),
        
        ("2. Vấn đề 'Trình diễn Công thức Toán' trong File Word (LaTeX to OMML)",
         "• Thắc mắc đã đặt ra: 'Làm sao khi kết quả trả về là file Word nhưng trong đó các công thức toán học được đưa về dạng chuẩn Word chứ không bị vỡ ảnh hay hiển thị chuỗi LaTeX thô?'\n"
         "• Bóc tách nguyên nhân & Vấn đề đang mắc phải: Đã xử lý thành công bằng thuật toán 3 bước (LaTeX -> MathML -> OMML) trong docx_exporter.py. Tuy nhiên, vấn đề mắc phải hiện tại là mới chỉ hỗ trợ các công thức dòng/khối cơ bản; chưa hỗ trợ được các sơ đồ toán học hay hình vẽ hình học phẳng (TikZ / Geometry) phức tạp.\n"
         "• Hướng giải quyết đề xuất: Kết hợp thư viện matplotlib / tikz2image render hình vẽ thành ảnh PNG chèn song song vào file Word."),
        
        ("3. Vấn đề Slide PowerPoint thô sơ & Hạn chế của Marp CLI",
         "• Thắc mắc đã đặt ra: 'Tại sao khi để AI sinh ra slide ở dự án này, slide khi được sinh ra như không thể chỉnh sửa linh hoạt, bị gò bó như 1 bức ảnh?'\n"
         "• Bóc tách nguyên nhân & Vấn đề đang mắc phải: Mã nguồn hiện tại vẫn đang phụ thuộc vào công cụ Marp CLI. Marp CLI render slide theo dạng các khối tĩnh văn bản, dẫn đến Slide PPTX sinh ra có thẩm mỹ rất kém, mang phong cách ghi chú thô, không có visual element sinh động và không có hình ảnh minh họa thật (mới chỉ có text image_prompt).\n"
         "• Hướng giải quyết đề xuất: Loại bỏ Marp CLI ở Giai đoạn 2 và chuyển sang sử dụng thư viện python-pptx native. Thiết kế các Slide Master Template sẵn có để chèn text/image/shape vào PPTX thật linh hoạt."),
        
        ("4. Vấn đề Bề nổi của Logic Kiểm duyệt Quy chuẩn 5512",
         "• Thắc mắc đã đặt ra: 'Xử lý AI và logic 5512 là những gì? Thể hiện qua các dòng code như thế nào? Liệu code của tôi đã giải quyết triệt để hay chưa?'\n"
         "• Bóc tách nguyên nhân & Vấn đề đang mắc phải: Code hiện tại trong validation.py mới chỉ giải quyết BỀ NỔI (ĐÚNG KHUNG JSON). Logic kiểm duyệt 5512 hiện tại rất ngây thơ: chỉ đếm số lượng 4 hoạt động và check từ khóa thô sơ. Hệ thống hoàn toàn chưa đánh giá được tính đúng đắn tri thức sư phạm hay chất lượng bài giảng.\n"
         "• Hướng giải quyết đề xuất: Xây dựng một AI Evaluator Agent riêng biệt ở Giai đoạn 2 chuyên chấm điểm sư phạm và phát hiện lỗi ảo giác (hallucination) của LLM."),
        
        ("5. Vấn đề Hoàn toàn vắng bóng Giao diện Người dùng (No UI) & Phụ thuộc API độc tôn",
         "• Bóc tách nguyên nhân & Vấn đề đang mắc phải: Tác giả hiện tại phải gọi code bằng lệnh Terminal hoặc Swagger UI. Giáo viên hoàn toàn không thể sử dụng hệ thống nếu không có Web UI. Ngoài ra, việc phụ thuộc độc tôn vào 1 API Gemini làm ứng dụng có nguy cơ sập toàn bộ nếu API hết Quota hoặc bị nghẽn mạng.\n"
         "• Hướng giải quyết đề xuất: Xây dựng Frontend Web UI bằng React/Next.js và thiết lập cơ chế Fallback LLM (Gemini -> OpenAI GPT-4o / Claude 3.5).")
    ]
    
    for title, desc in issues_analyzed:
        p_it = doc.add_paragraph()
        r_it = p_it.add_run(title)
        r_it.font.bold = True
        r_it.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
        
        p_id = doc.add_paragraph()
        p_id.paragraph_format.left_indent = Inches(0.2)
        p_id.paragraph_format.line_spacing = 1.15
        p_id.add_run(desc)
        doc.add_paragraph()
        
    # -------------------------------------------------------------
    # SECTION IV: ĐÁNH GIÁ THỰC CHẤT: KẾT QUẢ HIỆN TẠI THỂ HIỆN ĐIỀU GÌ?
    # -------------------------------------------------------------
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("IV. ĐÁNH GIÁ THỰC CHẤT: KẾT QUẢ GIAI ĐOẠN 1 THỂ HIỆN ĐIỀU GÌ?")
    r_h4.font.bold = True
    r_h4.font.size = Pt(13)
    r_h4.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_ans = doc.add_paragraph()
    p_ans.paragraph_format.line_spacing = 1.15
    p_ans.add_run(
        "Nhìn nhận một cách chân thực, trung thực và sát với thực tế dự án hiện tại, kết quả Giai đoạn 1 "
    )
    r_ev_b = p_ans.add_run("THỰC CHẤT THỂ HIỆN 3 ĐIỀU:")
    r_ev_b.font.bold = True
    p_ans.add_run("\n\n")
    
    eval_points = [
        ("1. CHỈ THỂ HIỆN TÍNH KHẢ THI VỀ MẶT THUẬT TOÁN (Proof of Concept - POC):",
         "Kết quả chạy được script/API hiện tại chứng minh rằng giải thuật kết hợp AI + Structured Output + Document Exporters là KHẢ THI. Luồng dữ liệu có thể chảy tự động từ Yêu cầu -> Gemini -> JSON -> File Word và PowerPoint mà không bị gãy đoạn."),
        
        ("2. HOÀN TOÀN CHƯA THỂ HIỆN TÍNH SẴN SÀNG SỬ DỤNG (0% Production-Ready):",
         "Dự án tuyệt đối CHƯA THỂ coi là một sản phẩm phần mềm hoàn chỉnh. Việc thiếu Giao diện Web, thiếu Cơ sở dữ liệu, Slide PPTX xấu và Validator thô sơ khiến dự án hiện tại mới chỉ là một 'khung xương Backend thô', chưa thể đưa cho bất kỳ giáo viên nào sử dụng thực tế."),
        
        ("3. ĐÓNG VAI TRÒ LÀM NỀN TẢNG THỬ NGHIỆM ĐỂ BỨC PHÁ Ở GIAI ĐOẠN 2:",
         "Những gì đã làm được đóng vai trò 'làm móng' (Foundation): xác nhận mô hình dữ liệu đúng và thử nghiệm thuật toán xuất file thành công. Toàn bộ phần việc quyết định giá trị thật sự của đồ án (Giao diện Web UI, Thiết kế Slide đẹp chuyên nghiệp, Đánh giá nội dung Sư phạm sâu, CSDL persistent) đều nằm ở phía trước.")
    ]
    
    for ep_t, ep_d in eval_points:
        p_ep = doc.add_paragraph()
        r_e1 = p_ep.add_run(f"{ep_t}\n")
        r_e1.font.bold = True
        r_e1.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        
        p_ed = doc.add_paragraph()
        p_ed.paragraph_format.left_indent = Inches(0.25)
        p_ed.paragraph_format.line_spacing = 1.15
        p_ed.add_run(ep_d)
        doc.add_paragraph()
        
    # -------------------------------------------------------------
    # SECTION V: KHUYẾN NGHỊ LỘ TRÌNH THỰC HIỆN GIAI ĐOẠN 2
    # -------------------------------------------------------------
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("V. KHUYẾN NGHỊ LỘ TRÌNH KHẮC PHỤC DÀNH CHO TÁC GIẢ Ở GIAI ĐOẠN 2")
    r_h5.font.bold = True
    r_h5.font.size = Pt(13)
    r_h5.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.line_spacing = 1.15
    p_rec.add_run(
        "Để chuyển đổi thành công đồ án từ một 'Script POC Backend thô' thành một 'Sản phẩm phần mềm thực sự hoàn chỉnh', "
        "tác giả cần triển khai ngay 5 bước trong Giai đoạn 2:\n\n"
    )
    
    next_steps = [
        ("Bước 1: Thiết lập CSDL (Database Integration)", "Tích hợp PostgreSQL + SQLAlchemy ORM để quản lý Users, Lưu Kế hoạch bài dạy, Lưu Slide Decks và Trạng thái phiên làm việc."),
        ("Bước 2: Xây dựng Giao diện Web UI (React / Next.js)", "Phát triển Web Frontend: Form nhập bài dạy, Live Editor chỉnh sửa giáo án, Web Slide Preview và nút Tải file Word/PPTX."),
        ("Bước 3: Chuyển đổi Slide Engine từ Marp CLI sang python-pptx Native", "Tự thiết kế Slide Master Templates bằng python-pptx để Slide sinh ra đẹp mắt, sắc nét và dễ chỉnh sửa."),
        ("Bước 4: Nâng cấp AI Pedagogical Evaluator Agent", "Bổ sung LLM Evaluator đánh giá chiều sâu chất lượng sư phạm và kiểm tra tính chính xác của kiến thức."),
        ("Bước 5: Tích hợp AI Image Gen & Async SSE Streaming", "Tích hợp API Flux/DALL-E sinh ảnh minh họa thật cho Slide và phát luồng Server-Sent Events (SSE) hiển thị tiến độ thời gian thực.")
    ]
    
    for s_t, s_d in next_steps:
        p_st = doc.add_paragraph()
        r_st1 = p_st.add_run(f"• {s_t}\n")
        r_st1.font.bold = True
        r_st1.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
        
        p_sd = doc.add_paragraph()
        p_sd.paragraph_format.left_indent = Inches(0.25)
        p_sd.paragraph_format.line_spacing = 1.15
        p_sd.add_run(s_d)
        doc.add_paragraph()
        
    output_path = "/Users/nguyennhan18/Documents/Lesson-Plan-Generator/Final_RP_1.docx"
    doc.save(output_path)
    print(f"✅ Final Stage 1 Report saved successfully to {output_path}")

if __name__ == "__main__":
    create_final_rp1()
