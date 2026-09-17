import io
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from schemas.schemas import LessonPlanSchema
import latex2mathml.converter
import mathml2omml

# Lọc và xóa các ký tự rác / ký tự ẩn không hợp lệ trong văn bản do AI sinh ra (nhằm tránh lỗi khi xuất file Word)
def clean_xml_string(s: str) -> str:
    if not isinstance(s, str):
        s = str(s or "")
    return re.sub(r'[\x00-\x08\x0B\x0E-\x1F\x7F-\x84\x86-\x9F]', '', s)

# Tô màu nền cho ô trong table
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" />')
    tcPr.append(shd)
_LATEX_REPAIRT_RULES = [
    (re.compile(r'(?<!\\)\brac(?=\{)'), r'\\frac'),
    (re.compile(r'(?<!\\)\bext(?=\{)'), r'\\text'),
    (re.compile(r'(?<!\\)\bqrt(?=\{)'), r'\\sqrt'),
    (re.compile(r'(?<!\\)\boxed(?=\{)'), r'\\boxed'),
    (re.compile(r'(?<!\\)\begin(?=\{)'), r'\\begin'),
]
def repair_latex_code(latex_code: str) -> str:



    repaired = latex_code
    for pattern, replacement in _LATEX_REPAIRT_RULES:
        repaired = pattern.sub(replacement, repaired)
    repaired = re.sub(r'[\x09\x0A\x0D]', '', repaired)
    return repaired
def convert_latex_to_omml(latex_code: str):
    """
    Chuyển đổi chuỗi mã LaTeX sang đối tượng XML OMML (Công thức toán native trong Word)
    """
    try:
        # 1. Chuyển LaTeX sang MathML
        mathml_code = latex2mathml.converter.convert(latex_code)    
        # 2. Chuyển MathML sang OMML XML
        omml_str = mathml2omml.convert(mathml_code)
        # 3. Thêm khai báo namespace nếu chưa có
        if '<m:oMath>' in omml_str and 'xmlns:m=' not in omml_str:
            omml_str = omml_str.replace('<m:oMath>', f'<m:oMath {nsdecls("m")}>')
        elif '<m:oMathPara>' in omml_str and 'xmlns:m=' not in omml_str:
            omml_str = omml_str.replace('<m:oMathPara>', f'<m:oMathPara {nsdecls("m")}>')
        # 4. Parse sang Element XML của docx
        return parse_xml(omml_str)
    except Exception as e:
        print(f"Lỗi chuyển đổi LaTeX sang OMML ('{latex_code}'): {e}")
        return None

def add_text_with_latex(paragraph, text: str, prefix: str = ""):
    r"""
    Phân tích văn bản và chèn vào paragraph.
    Tự động nhận diện các đoạn mã LaTeX dạng \(...\), \[...\], $...$, $$...$$
    và chuyển thành Công thức toán chuẩn của Word (OMML Native Equation).
    """
    cleaned_text = clean_xml_string(text)
    
    if prefix:
        r_pre = paragraph.add_run(prefix)
        r_pre.font.name = 'Times New Roman'

    if not cleaned_text:
        return

    # Regex bóc tách các thẻ công thức LaTeX: \(...\), \[...\], $...$, $$...$$
    pattern = r'(\\\(.*?\x5c\)|\\\[.*?\\\]|(?:\$\$[^$]+\$\$)|(?:\$[^$]+\$))'
    parts = re.split(pattern, cleaned_text)

    for part in parts:
        if not part:
            continue
        
        is_latex = False
        latex_code = ""
        if (part.startswith(r"\(") and part.endswith(r"\)")) or (part.startswith(r"\[") and part.endswith(r"\]")):
            is_latex = True
            latex_code = part[2:-2].strip()
        elif part.startswith("$$") and part.endswith("$$"):
            is_latex = True
            latex_code = part[2:-2].strip()
        elif part.startswith("$") and part.endswith("$"):
            is_latex = True
            latex_code = part[1:-1].strip()

        if is_latex and latex_code:
            omml_node = convert_latex_to_omml(latex_code)
            if omml_node is not None:
                paragraph._p.append(omml_node)
            else:
                # Fallback nếu convert lỗi: gán font Cambria Math
                run = paragraph.add_run(part)
                run.font.name = 'Cambria Math'
                run.font.italic = True
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'

def export_lesson_plan_to_docx(plan: LessonPlanSchema) -> io.BytesIO:
    doc = Document()
    # 1. Căn chỉnh lề (Margin)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.59)

    # 2. Định dạng font mặc định (Times New Roman, cỡ 13)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 3. Định dạng Header (Title)
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_header = p_header.add_run(f"GIÁO ÁN MÔN : {clean_xml_string(plan.subject).upper()} - LỚP {plan.grade}\n")
    run_header.font.bold = True
    run_header.font.size = Pt(14)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(f"  {clean_xml_string(plan.lesson_title).upper()}\n")
    run_title.font.bold = True
    run_title.font.size = Pt(15)

    # 4. Phần I: MỤC TIÊU BÀI HỌC
    h1 = doc.add_paragraph()
    r1 = h1.add_run("I. MỤC TIÊU BÀI HỌC")
    r1.font.bold = True

    if plan.knowledge_goals:
        p_k = doc.add_paragraph()
        r_k = p_k.add_run("1. Về kiến thức: ")
        r_k.font.bold = True
        for k in plan.knowledge_goals:
            p_item = doc.add_paragraph()
            add_text_with_latex(p_item, k, prefix=" - ")

    competencies = plan.competency_goals or plan.skills_goals
    if competencies:
        p_sk = doc.add_paragraph()
        r_sk = p_sk.add_run("2. Về năng lực: ")
        r_sk.font.bold = True
        for s in competencies:
            p_item = doc.add_paragraph()
            add_text_with_latex(p_item, s, prefix=" - ")

    if plan.character_goals:
        p_ch = doc.add_paragraph()
        r_ch = p_ch.add_run("3. Về phẩm chất: ")
        r_ch.font.bold = True
        for c in plan.character_goals:
            p_item = doc.add_paragraph()
            add_text_with_latex(p_item, c, prefix=" - ")

    # 5. Phần II: THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU
    if plan.teaching_equipment:
        h_eq = doc.add_paragraph()
        r_eq = h_eq.add_run("\nII. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU")
        r_eq.font.bold = True
        for eq in plan.teaching_equipment:
            p_eq_item = doc.add_paragraph()
            add_text_with_latex(p_eq_item, eq, prefix=" - ")

    # 6. Phần III: TIẾN TRÌNH DẠY HỌC
    h2 = doc.add_paragraph()
    r2 = h2.add_run("\nIII. TIẾN TRÌNH DẠY HỌC")
    r2.font.bold = True

    for act in plan.activities:
        # Tạo tiêu đề hoạt động
        p_act = doc.add_paragraph()
        act_goal = getattr(act, "goal", f"HOẠT ĐỘNG {act.activity_number}")
        r_act = p_act.add_run(f"\n HOẠT ĐỘNG {act.activity_number} : {clean_xml_string(act_goal).upper()} ({act.time_minutes} phút)")
        r_act.font.bold = True

        p_goal = doc.add_paragraph()
        add_text_with_latex(p_goal, act_goal, prefix="a) Mục tiêu: ")

        p_product = doc.add_paragraph()
        add_text_with_latex(p_product, act.expected_product, prefix="b) Sản phẩm dự kiến: ")

        p_exec = doc.add_paragraph()
        r_exec = p_exec.add_run("c) Tổ chức thực hiện: ")
        r_exec.font.bold = True

        # Vẽ bảng 4 bước tổ chức
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header bảng
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Các bước tổ chức"
        hdr_cells[1].text = "Nội dung chi tiết"
        set_cell_background(hdr_cells[0], "E0E0E0")
        set_cell_background(hdr_cells[1], "E0E0E0")
        if hdr_cells[0].paragraphs[0].runs:
            hdr_cells[0].paragraphs[0].runs[0].font.bold = True
        if hdr_cells[1].paragraphs[0].runs:
            hdr_cells[1].paragraphs[0].runs[0].font.bold = True
        
        # Data 4 bước
        steps_data = [
            ("Bước 1: Chuyển giao nhiệm vụ", act.execution.step_1_assign),
            ("Bước 2: Thực hiện nhiệm vụ", act.execution.step_2_excute),
            ("Bước 3: Báo cáo, kết luận", act.execution.step_3_report),
            ("Bước 4: Kết luận & nhận định", act.execution.step_4_conclusion),
        ]

        for idx, (step_name, step_content) in enumerate(steps_data, start=1):
            row_cells = table.rows[idx].cells
            row_cells[0].text = step_name
            if row_cells[0].paragraphs[0].runs:
                row_cells[0].paragraphs[0].runs[0].font.bold = True
            
            p_content = row_cells[1].paragraphs[0]
            p_content.text = ""  # Xóa text mặc định
            add_text_with_latex(p_content, step_content)

    # Trả kết quả
    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream