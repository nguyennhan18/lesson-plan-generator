import os
import sys
import re
import io
import json
import math
from enum import Enum
from typing import Union, Dict, Any, List
import matplotlib.pyplot as plt
import sympy as sp
import numpy as np

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from PIL import Image

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from schemas.slide_schema import SlideDeckSchema, SlideType, ThemeType
from exporters.graph_generator import generate_math_plot, is_plottable


def clean_xml_string(s: str) -> str:
    """Lọc ký tự rác / ẩn để tránh lỗi khi chèn XML vào PowerPoint"""
    if not isinstance(s, str):
        s = str(s or "")
    return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x84\x86-\x9F]', '', s)


_GREEK_MAP = {
    r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\delta': 'δ',
    r'\epsilon': 'ε', r'\varepsilon': 'ε', r'\zeta': 'ζ', r'\eta': 'η',
    r'\theta': 'θ', r'\vartheta': 'ϑ', r'\iota': 'ι', r'\kappa': 'κ',
    r'\lambda': 'λ', r'\mu': 'μ', r'\nu': 'ν', r'\xi': 'ξ',
    r'\pi': 'π', r'\rho': 'ρ', r'\sigma': 'σ', r'\tau': 'τ',
    r'\upsilon': 'υ', r'\phi': 'φ', r'\varphi': 'φ', r'\chi': 'χ',
    r'\psi': 'ψ', r'\omega': 'ω',
    r'\Gamma': 'Γ', r'\Delta': 'Δ', r'\Theta': 'Θ', r'\Lambda': 'Λ',
    r'\Xi': 'Ξ', r'\Pi': 'Π', r'\Sigma': 'Σ', r'\Upsilon': 'Υ',
    r'\Phi': 'Φ', r'\Psi': 'Ψ', r'\Omega': 'Ω',
}

_SYMBOL_MAP = {
    r'\times': '×', r'\div': '÷', r'\cdot': '·', r'\pm': '±', r'\mp': '∓',
    r'\leq': '≤', r'\geq': '≥', r'\neq': '≠', r'\approx': '≈', r'\equiv': '≡',
    r'\infty': '∞', r'\partial': '∂', r'\nabla': '∇', r'\emptyset': '∅',
    r'\forall': '∀', r'\exists': '∃', r'\in': '∈', r'\notin': '∉',
    r'\subset': '⊂', r'\supset': '⊃', r'\subseteq': '⊆', r'\supseteq': '⊇',
    r'\cup': '∪', r'\cap': '∩', r'\Rightarrow': '⇒', r'\Leftarrow': '⇐',
    r'\Leftrightarrow': '⇔', r'\rightarrow': '→', r'\leftarrow': '←',
    r'\leftrightarrow': '↔', r'\to': '→', r'\mathbb{R}': 'ℝ', r'\mathbb{N}': 'ℕ',
    r'\mathbb{Z}': 'ℤ', r'\mathbb{Q}': 'ℚ', r'\mathbb{C}': 'ℂ',
    r'\sqrt': '√', r'\int': '∫', r'\sum': '∑', r'\prod': '∏',
    r'\cos': 'cos', r'\sin': 'sin', r'\tan': 'tan', r'\cot': 'cot', r'\lim': 'lim',
    r'\log': 'log', r'\ln': 'ln',
}

_SUPERSCRIPT_MAP = {
    '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
    '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
    '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
    'n': 'ⁿ', 'i': 'ⁱ', 'x': 'ˣ', 'y': 'ʸ', 'k': 'ᵏ', 'm': 'ᵐ'
}

_SUBSCRIPT_MAP = {
    '0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
    '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉',
    '+': '₊', '-': '₋', '=': '₌', '(': '₍', ')': '₎',
    'a': 'ₐ', 'e': 'ₑ', 'h': 'ₕ', 'i': 'ᵢ', 'j': 'ⱼ',
    'k': 'ₖ', 'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ', 'o': 'ₒ',
    'p': 'ₚ', 'r': 'ᵣ', 's': 'ₛ', 't': 'ₜ', 'u': 'ᵤ',
    'v': 'ᵥ', 'x': 'ₓ'
}


def clean_latex_table_environment(text: str) -> str:
    """Biến đổi môi trường array/matrix/table của LaTeX thành dạng bảng văn bản thuần dòng/cột."""
    def replace_array(match):
        content = match.group(1)
        content = re.sub(r'\\hline|\\vline', '', content)
        rows = content.split(r'\\')
        cleaned_rows = []
        for r in rows:
            r_str = r.strip()
            if not r_str:
                continue
            cols = [c.strip() for c in r_str.split('&')]
            cleaned_rows.append(' | '.join(cols))
        return '\n' + '\n'.join(cleaned_rows) + '\n'

    pattern = r'\\begin\{(?:array|matrix|pmatrix|bmatrix|vmatrix)\}(?:\{.*?\})?(.*?)\\end\{(?:array|matrix|pmatrix|bmatrix|vmatrix)\}'
    return re.sub(pattern, replace_array, text, flags=re.DOTALL)


def latex_to_unicode_math(text: str) -> str:
    """
    Chuyển đổi mã LaTeX thành ký tự Unicode toán học chuẩn cho PowerPoint DrawingML.
    Áp dụng toàn diện các bước cho TOÀN BỘ chuỗi văn bản (không cần dấu $ bao quanh).
    """
    if not isinstance(text, str):
        return str(text or '')

    text = clean_xml_string(text)
    text = clean_latex_table_environment(text)
    text = text.replace(r'\(', '').replace(r'\)', '').replace('$', '')

    for k in sorted(_GREEK_MAP.keys(), key=len, reverse=True):
        text = text.replace(k, _GREEK_MAP[k])

    for k in sorted(_SYMBOL_MAP.keys(), key=len, reverse=True):
        text = text.replace(k, _SYMBOL_MAP[k])

    def convert_script(match):
        script_type, body = match.groups()
        char_map = _SUPERSCRIPT_MAP if script_type == '^' else _SUBSCRIPT_MAP
        return ''.join(char_map.get(ch, ch) for ch in body)

    text = re.sub(r'([\^_])\{([^}]+)\}', convert_script, text)
    text = re.sub(r'([\^_])([0-9a-zA-Z\+\-\=\(\)])', convert_script, text)

    text = re.sub(r'\\(?:newline|quad|qquad|left|right|text|mathrm|mathbf|mathit|vec|bar|hat)', ' ', text)
    text = text.replace('{', '').replace('}', '')
    text = re.sub(r'\\+', ' ', text)

    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()


def render_latex_to_bytes(latex_str: str) -> bytes:
    """
    Render công thức LaTeX ra ảnh PNG chuẩn nét (300 DPI).
    Tự động xử lý và fallback an toàn nếu gặp ký tự tiếng Việt hoặc TeX không được mathtext hỗ trợ.
    """
    clean_str = latex_str.replace(r'\(', '').replace(r'\)', '').strip()

    clean_str = re.sub(r'\\ge(?![a-zA-Z])', r'\\geq', clean_str)
    clean_str = re.sub(r'\\le(?![a-zA-Z])', r'\\leq', clean_str)
    clean_str = re.sub(r'\\implies(?![a-zA-Z])', r'\\Rightarrow', clean_str)
    clean_str = re.sub(r'\\iff(?![a-zA-Z])', r'\\Leftrightarrow', clean_str)
    clean_str = re.sub(r'\\to(?![a-zA-Z])', r'\\rightarrow', clean_str)
    clean_str = clean_str.replace(r'\uparrow', r'\nearrow').replace(r'\downarrow', r'\searrow')
    clean_str = clean_str.replace(r'\quad', ' ').replace(r'\qquad', ' ')
    clean_str = clean_str.replace(r'\tg', r'\tan').replace(r'\ctg', r'\cot')
    clean_str = re.sub(r'\\text\{([^}]+)\}', r'\1', clean_str)
    clean_str = re.sub(r'\\mathrm\{([^}]+)\}', r'\1', clean_str)

    fig, ax = plt.subplots(figsize=(0.1, 0.1))
    ax.axis('off')

    try:
        formatted_latex = f"${clean_str}$" if not clean_str.startswith("$") else clean_str
        ax.text(
            0.5, 0.5, formatted_latex,
            fontsize=20, ha='center', va='center', color='#1e3a8a'
        )
        buf = io.BytesIO()
        plt.savefig(buf, format='png', bbox_inches='tight', pad_inches=0.04, dpi=300, transparent=True)
        plt.close(fig)
        buf.seek(0)
        return buf.getvalue()
    except Exception:
        plt.close(fig)
        # Fallback render bằng Unicode text
        fig, ax = plt.subplots(figsize=(0.1, 0.1))
        ax.axis('off')
        unicode_str = latex_to_unicode_math(latex_str)
        ax.text(
            0.5, 0.5, unicode_str,
            fontsize=18, ha='center', va='center', color='#1e3a8a', fontname='Arial'
        )
        buf = io.BytesIO()
        plt.savefig(buf, format='png', bbox_inches='tight', pad_inches=0.04, dpi=300, transparent=True)
        plt.close(fig)
        buf.seek(0)
        return buf.getvalue()


def _plain(obj):
    """Đưa Enum / Pydantic model về kiểu Python thuần (str, dict, list)."""
    if isinstance(obj, Enum):
        return obj.value
    if hasattr(obj, "model_dump"):
        return _plain(obj.model_dump())
    if isinstance(obj, dict):
        return {k: _plain(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_plain(v) for v in obj]
    return obj


# Bảng màu chuẩn Academic Sư phạm (Đồng bộ 100% với giao diện Studio Web)
COLOR_NAVY = RGBColor(15, 43, 92)          # #0f2b5c: Header & tiêu đề chính
COLOR_ROYAL_BLUE = RGBColor(37, 99, 235)   # #2563eb: Điểm nhấn, nút, bullets
COLOR_ACCENT_BLUE = RGBColor(59, 130, 246) # #3b82f6: Đường viền accent header
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)     # #1e293b: Văn bản nội dung
COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # #475569: Phụ đề & chú thích
COLOR_FOOTER = RGBColor(148, 163, 184)     # #94a3b8: Footer trang
COLOR_BG_FORMULA = RGBColor(248, 250, 252) # #f8fafc: Khung công thức
COLOR_BORDER_FORMULA = RGBColor(191, 219, 254) # #bfdbfe: Viền khung công thức
COLOR_CARD_BORDER = RGBColor(226, 232, 240)    # #e2e8f0: Viền card đồ thị
COLOR_WHITE = RGBColor(255, 255, 255)


def build_slide_header(slide, slide_idx: int, total_slides: int, slide_type: str):
    """Vẽ thanh Header Academic Navy ở đỉnh slide (Mô phỏng chuẩn Web Studio)"""
    # 1. Thanh nền Navy
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10.0), Inches(0.65))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_NAVY
    top_bar.line.fill.background()

    # 2. Đường kẻ accent xanh dương
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.65), Inches(10.0), Inches(0.04))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    accent.line.fill.background()

    # 3. Badge số thứ tự slide (bên trái)
    badge_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.13), Inches(1.3), Inches(0.38))
    badge_l.fill.solid()
    badge_l.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge_l.line.color.rgb = RGBColor(96, 165, 250)
    badge_l.line.width = Pt(1)
    tf_l = badge_l.text_frame
    tf_l.word_wrap = False
    tf_l.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_l = tf_l.paragraphs[0]
    p_l.alignment = PP_ALIGN.CENTER
    r_l = p_l.add_run()
    r_l.text = f'SLIDE {slide_idx}/{total_slides}'
    r_l.font.name = 'Arial'
    r_l.font.size = Pt(9.5)
    r_l.font.bold = True
    r_l.font.color.rgb = COLOR_WHITE

    # 4. Badge loại slide (bên phải)
    clean_type = str(slide_type or 'SLIDE').replace('_', ' ').upper()
    badge_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(0.13), Inches(1.5), Inches(0.38))
    badge_r.fill.solid()
    badge_r.fill.fore_color.rgb = COLOR_ROYAL_BLUE
    badge_r.line.fill.background()
    tf_r = badge_r.text_frame
    tf_r.word_wrap = False
    tf_r.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.CENTER
    r_r = p_r.add_run()
    r_r.text = clean_type
    r_r.font.name = 'Arial'
    r_r.font.size = Pt(8.5)
    r_r.font.bold = True
    r_r.font.color.rgb = COLOR_WHITE


def build_title_slide(slide, slide_data: dict, subject: str, grade: Any):
    """Thiết kế Slide Bìa (TITLE_SLIDE) chuẩn Sư phạm trang nhã"""
    # Tag Pill Môn học & Lớp
    tag_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(1.15), Inches(3.6), Inches(0.42))
    tag_pill.fill.solid()
    tag_pill.fill.fore_color.rgb = RGBColor(239, 246, 255)
    tag_pill.line.color.rgb = COLOR_BORDER_FORMULA
    tag_pill.line.width = Pt(1)
    tf_pill = tag_pill.text_frame
    tf_pill.word_wrap = False
    tf_pill.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_pill = tf_pill.paragraphs[0]
    p_pill.alignment = PP_ALIGN.CENTER
    r_pill = p_pill.add_run()
    r_pill.text = f'MÔN {str(subject).upper()} - LỚP {grade}'
    r_pill.font.name = 'Arial'
    r_pill.font.size = Pt(10.5)
    r_pill.font.bold = True
    r_pill.font.color.rgb = COLOR_ROYAL_BLUE

    # Tiêu đề bài học chính (Hero Title)
    title_text = latex_to_unicode_math(slide_data.get('title') or '')
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(1.4))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.CENTER
    r_t = p_title.add_run()
    r_t.text = title_text
    r_t.font.name = 'Times New Roman'
    font_pt = 32 if len(title_text) < 40 else (28 if len(title_text) < 65 else 24)
    r_t.font.size = Pt(font_pt)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY

    # Phụ đề bài học
    subtitle_text = latex_to_unicode_math(slide_data.get('subtitle') or '')
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.3), Inches(8.0), Inches(0.9))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.alignment = PP_ALIGN.CENTER
        r_sub = p_sub.add_run()
        r_sub.text = subtitle_text
        r_sub.font.name = 'Times New Roman'
        r_sub.font.size = Pt(16)
        r_sub.font.italic = True
        r_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Tagline chuẩn Công văn 5512
    tag_box = slide.shapes.add_textbox(Inches(2.0), Inches(4.55), Inches(6.0), Inches(0.4))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    r_tag = p_tag.add_run()
    r_tag.text = 'KẾ HOẠCH BÀI DẠY THEO CÔNG VĂN 5512'
    r_tag.font.name = 'Arial'
    r_tag.font.size = Pt(10)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(100, 116, 139)


def build_content_slide(slide, slide_data: dict, slide_idx: int, total_slides: int, subject: str, grade: Any):
    """Thiết kế Slide Nội dung (CONCEPT, EXERCISE, SUMMARY, v.v.)"""
    # 1. Tiêu đề slide
    title_text = latex_to_unicode_math(slide_data.get('title') or '')
    t_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.8), Inches(0.55))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = title_text
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY

    # 2. Phụ đề slide (nếu có)
    subtitle_text = latex_to_unicode_math(slide_data.get('subtitle') or '')
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(8.8), Inches(0.35))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        r_sub = p_sub.add_run()
        r_sub.text = subtitle_text
        r_sub.font.name = 'Times New Roman'
        r_sub.font.size = Pt(13)
        r_sub.font.italic = True
        r_sub.font.color.rgb = COLOR_TEXT_MUTED
        content_top = Inches(1.75)
    else:
        content_top = Inches(1.45)

    chart_spec = slide_data.get('chart_spec')
    has_chart = is_plottable(chart_spec)

    bullets = [b for b in (slide_data.get('bullet_points') or []) if b]
    formula_latex = slide_data.get('main_definition_latex') or ''

    # TRƯỜNG HỢP 1: CÓ ĐỒ THỊ HÀM SỐ (BỐ CỤC 2 CỘT CHUẨN WEB STUDIO)
    if has_chart:
        w_left = Inches(4.8)

        # Cột Trái: Bullet Points
        text_h = Inches(1.9) if formula_latex else Inches(3.2)
        text_box = slide.shapes.add_textbox(Inches(0.6), content_top, w_left, text_h)
        tf_c = text_box.text_frame
        tf_c.word_wrap = True
        for b_idx, b in enumerate(bullets):
            p_b = tf_c.paragraphs[0] if b_idx == 0 else tf_c.add_paragraph()
            p_b.space_after = Pt(4)
            r_bullet = p_b.add_run()
            r_bullet.text = '•  '
            r_bullet.font.name = 'Arial'
            r_bullet.font.bold = True
            r_bullet.font.size = Pt(13)
            r_bullet.font.color.rgb = COLOR_ROYAL_BLUE
            r_txt = p_b.add_run()
            r_txt.text = latex_to_unicode_math(b)
            r_txt.font.name = 'Times New Roman'
            r_txt.font.size = Pt(13)
            r_txt.font.color.rgb = COLOR_TEXT_MAIN

        # Cột Trái: Formula Card Box (nằm phía dưới bullets)
        if formula_latex:
            f_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3.85), w_left, Inches(1.2))
            f_box.fill.solid()
            f_box.fill.fore_color.rgb = COLOR_BG_FORMULA
            f_box.line.color.rgb = COLOR_BORDER_FORMULA
            f_box.line.width = Pt(1)

            f_accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3.85), Inches(0.06), Inches(1.2))
            f_accent.fill.solid()
            f_accent.fill.fore_color.rgb = COLOR_ROYAL_BLUE
            f_accent.line.fill.background()

            try:
                png_bytes = render_latex_to_bytes(formula_latex)
                with Image.open(io.BytesIO(png_bytes)) as im:
                    px_w, px_h = im.size
                asp = px_w / max(px_h, 1)
                img_h = Inches(0.7)
                img_w = Inches(0.7 * asp)
                if img_w > Inches(4.4):
                    img_w = Inches(4.4)
                    img_h = Inches(4.4 / asp)
                img_left = Inches(0.6) + (w_left - img_w) / 2
                img_top = Inches(3.85) + (Inches(1.2) - img_h) / 2
                slide.shapes.add_picture(io.BytesIO(png_bytes), img_left, img_top, width=img_w, height=img_h)
            except Exception as e:
                print(f"Lỗi chèn ảnh công thức: {e}")

        # Cột Phải: Matplotlib Plot Card
        card_right = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.6), content_top, Inches(3.8), Inches(3.4))
        card_right.fill.solid()
        card_right.fill.fore_color.rgb = COLOR_WHITE
        card_right.line.color.rgb = COLOR_CARD_BORDER
        card_right.line.width = Pt(1)

        try:
            chart_png = generate_math_plot(chart_spec, figsize=(4.8, 3.4))
            slide.shapes.add_picture(io.BytesIO(chart_png), Inches(5.7), content_top + Inches(0.08), width=Inches(3.6), height=Inches(2.45))
        except Exception as e:
            print(f"Lỗi chèn ảnh đồ thị: {e}")

        # Badge Đồ thị Matplotlib Native
        p_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.45), content_top + Inches(2.62), Inches(2.1), Inches(0.26))
        p_badge.fill.solid()
        p_badge.fill.fore_color.rgb = RGBColor(224, 242, 254)
        p_badge.line.fill.background()
        tf_pb = p_badge.text_frame
        tf_pb.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_run = tf_pb.paragraphs[0].add_run()
        p_run.text = '📊 Đồ thị Matplotlib Native'
        p_run.font.name = 'Arial'
        p_run.font.size = Pt(8.5)
        p_run.font.bold = True
        p_run.font.color.rgb = RGBColor(3, 105, 161)

        if chart_spec.get('chart_title'):
            ct_box = slide.shapes.add_textbox(Inches(5.7), content_top + Inches(2.9), Inches(3.6), Inches(0.35))
            ct_p = ct_box.text_frame.paragraphs[0]
            ct_p.alignment = PP_ALIGN.CENTER
            ct_r = ct_p.add_run()
            ct_r.text = chart_spec.get('chart_title')
            ct_r.font.name = 'Arial'
            ct_r.font.size = Pt(8.5)
            ct_r.font.color.rgb = COLOR_TEXT_MUTED

    # TRƯỜNG HỢP 2: KHÔNG CÓ ĐỒ THỊ (BỐ CỤC TOÀN CHIỀU RỘNG TRANG NHÃ)
    else:
        w_full = Inches(8.8)
        text_h = Inches(2.0) if formula_latex else Inches(3.3)
        text_box = slide.shapes.add_textbox(Inches(0.6), content_top, w_full, text_h)
        tf_c = text_box.text_frame
        tf_c.word_wrap = True
        font_pt = 15 if len(bullets) > 3 else 16
        for b_idx, b in enumerate(bullets):
            p_b = tf_c.paragraphs[0] if b_idx == 0 else tf_c.add_paragraph()
            p_b.space_after = Pt(6)
            r_bullet = p_b.add_run()
            r_bullet.text = '•  '
            r_bullet.font.name = 'Arial'
            r_bullet.font.bold = True
            r_bullet.font.size = Pt(font_pt)
            r_bullet.font.color.rgb = COLOR_ROYAL_BLUE
            r_txt = p_b.add_run()
            r_txt.text = latex_to_unicode_math(b)
            r_txt.font.name = 'Times New Roman'
            r_txt.font.size = Pt(font_pt)
            r_txt.font.color.rgb = COLOR_TEXT_MAIN

        if formula_latex:
            f_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.85), Inches(7.0), Inches(1.2))
            f_box.fill.solid()
            f_box.fill.fore_color.rgb = COLOR_BG_FORMULA
            f_box.line.color.rgb = COLOR_BORDER_FORMULA
            f_box.line.width = Pt(1)

            f_accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.85), Inches(0.06), Inches(1.2))
            f_accent.fill.solid()
            f_accent.fill.fore_color.rgb = COLOR_ROYAL_BLUE
            f_accent.line.fill.background()

            try:
                png_bytes = render_latex_to_bytes(formula_latex)
                with Image.open(io.BytesIO(png_bytes)) as im:
                    px_w, px_h = im.size
                asp = px_w / max(px_h, 1)
                img_h = Inches(0.8)
                img_w = Inches(0.8 * asp)
                if img_w > Inches(6.4):
                    img_w = Inches(6.4)
                    img_h = Inches(6.4 / asp)
                img_left = Inches(1.5) + (Inches(7.0) - img_w) / 2
                img_top = Inches(3.85) + (Inches(1.2) - img_h) / 2
                slide.shapes.add_picture(io.BytesIO(png_bytes), img_left, img_top, width=img_w, height=img_h)
            except Exception as e:
                print(f"Lỗi chèn ảnh công thức: {e}")

    # 3. Footer tinh tế ở chân slide
    ft_l = slide.shapes.add_textbox(Inches(0.6), Inches(5.26), Inches(5.0), Inches(0.25))
    p_ftl = ft_l.text_frame.paragraphs[0]
    r_ftl = p_ftl.add_run()
    r_ftl.text = f'Kế hoạch bài dạy môn {subject} - Lớp {grade}'
    r_ftl.font.name = 'Arial'
    r_ftl.font.size = Pt(8.5)
    r_ftl.font.color.rgb = COLOR_FOOTER

    ft_r = slide.shapes.add_textbox(Inches(8.0), Inches(5.26), Inches(1.4), Inches(0.25))
    p_ftr = ft_r.text_frame.paragraphs[0]
    p_ftr.alignment = PP_ALIGN.RIGHT
    r_ftr = p_ftr.add_run()
    r_ftr.text = f'Trang {slide_idx}/{total_slides}'
    r_ftr.font.name = 'Arial'
    r_ftr.font.size = Pt(8.5)
    r_ftr.font.color.rgb = COLOR_FOOTER


def export_slide_deck_to_pptx(data: Union[Dict[str, Any], SlideDeckSchema]) -> io.BytesIO:
    """
    Chuyển đổi dữ liệu SlideDeck thành file PowerPoint (.pptx) chuẩn 16:9 Widescreen.
    Đồng bộ 100% về bố cục, màu sắc, font chữ và đồ thị Matplotlib với giao diện Web Studio.
    Cho phép Giáo viên chỉnh sửa, thay đổi nội dung, đổi màu tự do 100% trong PowerPoint.
    """
    if isinstance(data, (SlideDeckSchema, dict)):
        deck_dict = _plain(data)
    else:
        raise ValueError("Dữ liệu vào phải là dict hoặc SlideDeckSchema")

    slides_data = deck_dict.get("slides", [])
    total_slides = len(slides_data)
    subject = deck_dict.get("subject", "Toán học")
    grade = deck_dict.get("grade", 12)

    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(5.625)

    for idx, slide_data in enumerate(slides_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Xóa các placeholder thừa nếu có
        for sh in list(slide.placeholders):
            slide.shapes._spTree.remove(sh._element)

        slide_type = str(slide_data.get("type") or "CONCEPT_SLIDE")
        is_title = (slide_type == "TITLE_SLIDE")

        # 1. Header Banner
        build_slide_header(slide, idx + 1, total_slides, slide_type)

        # 2. Nội dung slide
        if is_title:
            build_title_slide(slide, slide_data, subject, grade)
        else:
            build_content_slide(slide, slide_data, idx + 1, total_slides, subject, grade)

        # 3. Kịch bản lời thoại giáo viên (Speaker Notes Pane trong PowerPoint)
        teacher_note = slide_data.get("teacher_note", "")
        if teacher_note:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = f"[Kịch bản giảng dạy dành cho Giáo viên]:\n{teacher_note}"

    pptx_bytes = io.BytesIO()
    prs.save(pptx_bytes)
    pptx_bytes.seek(0)
    return pptx_bytes


def export_slide_deck_to_pptx_file(data: Union[Dict[str, Any], SlideDeckSchema], output_filepath: str) -> str:
    """Xuất bài giảng ra file .pptx trên đĩa cứng"""
    pptx_stream = export_slide_deck_to_pptx(data)
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    with open(output_filepath, "wb") as f:
        f.write(pptx_stream.getbuffer())
    print(f"Đã xuất thành công file PowerPoint PPTX đồng bộ chuẩn Web Studio: {output_filepath}")
    return output_filepath


# Alias tương thích ngược
export_slide_to_pptx = export_slide_deck_to_pptx_file


def render_math_plot(chart_spec: dict) -> io.BytesIO:
    """Vẽ đồ thị toán học từ cấu hình chart_spec và trả về buffer ảnh (BytesIO)."""
    return io.BytesIO(generate_math_plot(chart_spec))


def load_manifest_config() -> dict:
    """Đọc file manifest.json trong thư mục templates/ (giữ cho tương thích ngược)"""
    app_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    manifest_path = os.path.join(app_dir, "templates", "manifest.json")
    if os.path.exists(manifest_path):
        with open(manifest_path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def select_template_presentation(theme_id: str = "math_academic") -> tuple[Presentation, dict]:
    """Tạo Presentation 16:9 widescreen đồng bộ với giao diện web"""
    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(5.625)
    return prs, {}
