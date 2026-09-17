from pptx import Presentation
import pptx
from pptx import presentation
from pptx.dml.color import RGBColor

def create_pptx_from_json(slide_deck_data: dict, output_filepath: str):
    prs = Presentation()
    
    for slide_data in slide_deck_data["slides"]:
        slide_type = slide_data.get("type")
        title_text = slide_data.get("title", "")
        teacher_note = slide_data.get("teacher_note", "")
        bullet_points = slide_data.get("bullet_points", [])

        if slide_type == "TITLE_SLIDE":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            title_box = slide.shapes.title
            subtitle_box = slide.placeholders[1]

            title_box.text = title_text
            subtitle_box.text = slide_data.get("subtitle" , "")

        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])

            slide.shapes.title.text = title_text

            tf = slide.placeholders[1].text_frame
            tf.word_wrap = True

            for idx, point in enumerate(bullet_points):
                if idx == 0:
                    p = tf.paragraphs[0]
                    p.text = point
                else:
                    p = tf.add_paragraph()
                    p.text = point        
                
            if slide_data.get("main_definition_latex"):
                latex_str = slide_data["main_definition_latex"]
                p = tf.add_paragraph()
                p.text = f"Cong thuc: {latex_str}"
                p.font.color = True
                p.font.color.rgb = RGBColor(0, 102, 204)

        if teacher_note:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = f"[Loi thoai GV] : {teacher_note}"
    prs.save(output_filepath)
    print(f"Da hoan thanh file PPT: {output_filepath}")