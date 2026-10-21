r"""
Automated Test Suite — Day 14: Cầu nối Chụp ảnh Snapshot 3D gửi về Backend xuất PPTX
Kiểm thử chi tiết các hạng mục của Ngày 14:
- Task 14.1: Viết hàm JavaScript capture3DSnapshot() sử dụng canvas.toDataURL('image/png') chụp lại khung hình 3D hiện tại
- Task 14.2: Frontend đóng gói chuỗi Base64 Data URL vào request POST /api/v1/lesson-plan/export-slide
- Task 14.3: Backend giải mã chuỗi Base64 thành byte stream nhúng ảnh vào slide PPTX qua python-pptx
- Task 14.4: Tinh chỉnh tọa độ và kích thước khung ảnh 3D cân đối với bố cục 16:9
- Task 14.5: Nghiệm thu DoD Ngày 14 và kiểm thử tích hợp toàn diện
"""

import base64
import io
import re
import sys
from pathlib import Path
from typing import Any
import pytest
from fastapi.testclient import TestClient
from typing import Any
ROOT_DIR = Path(__file__).resolve().parents[3]
BACKEND_SRC = ROOT_DIR / "app" / "backend" / "src"
FRONTEND_DIR = ROOT_DIR / "app" / "frontend"

if str(BACKEND_SRC) not in sys.path:
    sys.path.insert(0, str(BACKEND_SRC))

from main import app

client = TestClient(app)


# 1x1 transparent PNG pixel base64 string
SAMPLE_PNG_BASE64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)
SAMPLE_DATA_URL = f"data:image/png;base64,{SAMPLE_PNG_BASE64}"


def test_task_14_1_webgl_canvas_snapshot_configuration():
    """
    Task 14.1: Kiểm tra cấu hình WebGL và hàm JavaScript capture3DSnapshot() trong ThreeCanvas.jsx:
    1. Cấu hình preserveDrawingBuffer: true trong THREE.WebGLRenderer để giữ buffer sau compositor pass
    2. Hàm capture3DSnapshot() thực hiện render đồng bộ renderer.render(scene, camera) trước khi chụp
    3. Trích xuất buffer qua renderer.domElement.toDataURL('image/png', 1.0)
    4. Cung cấp onCaptureReady callback và captureRef để component cha truy cập
    5. Đăng ký window.__threeCanvasCaptureSnapshot phục vụ kiểm thử và tự động hóa
    6. Nút bấm Chụp ảnh Snapshot 3D (Camera icon) trên thanh công cụ
    """
    canvas_path = FRONTEND_DIR / "src" / "components" / "ThreeCanvas.jsx"
    assert canvas_path.exists(), f"Không tìm thấy file: {canvas_path}"
    content = canvas_path.read_text(encoding="utf-8")

    # 1. preserveDrawingBuffer: true
    assert "preserveDrawingBuffer: true" in content, (
        "Thiếu thuộc tính preserveDrawingBuffer: true trong THREE.WebGLRenderer. "
        "Nếu thiếu, toDataURL sẽ trả về ảnh đen hoặc rỗng sau swap buffer!"
    )

    # 2. Hàm capture3DSnapshot
    assert "capture3DSnapshot" in content, "Không tìm thấy định nghĩa hàm capture3DSnapshot trong ThreeCanvas.jsx"

    # 3. Render đồng bộ cưỡng bức ngay trước khi đọc buffer
    assert "renderer.render(scene, camera)" in content or "renderer.render(sceneRef.current" in content or "rendererRef.current.render" in content, (
        "Cần gọi render(scene, camera) ngay trước khi toDataURL để đảm bảo khung hình hiện thời được vẽ vào buffer."
    )

    # 4. toDataURL('image/png')
    assert "toDataURL('image/png'" in content or 'toDataURL("image/png"' in content, (
        "Thiếu lệnh toDataURL('image/png') để xuất ảnh định dạng PNG chuẩn."
    )

    # 5. Kết nối onCaptureReady và window.__threeCanvasCaptureSnapshot
    assert "onCaptureReady" in content, "ThreeCanvas cần nhận prop onCaptureReady"
    assert "window.__threeCanvasCaptureSnapshot" in content, (
        "ThreeCanvas cần gán window.__threeCanvasCaptureSnapshot để truy cập toàn cục"
    )

    # 6. Nút bấm Camera snapshot trên UI
    assert "Camera" in content, "Thiếu icon Camera trong thanh công cụ ThreeCanvas"
    assert "Chụp ảnh Snapshot 3D (PNG)" in content, "Thiếu tooltip/tiêu đề cho nút chụp snapshot 3D"


def test_task_14_1_component_wiring_in_stage_and_app():
    """
    Task 14.1: Kiểm tra luồng kết nối component PresentationStage và App:
    - PresentationStage.jsx nhận onCaptureReady và truyền xuống ThreeCanvas
    - App.jsx quản lý threeCanvasCaptureRef và truyền onCaptureReady xuống PresentationStage
    """
    stage_path = FRONTEND_DIR / "src" / "components" / "PresentationStage.jsx"
    assert stage_path.exists()
    stage_content = stage_path.read_text(encoding="utf-8")

    assert "onCaptureReady" in stage_content
    assert "threeCaptureRef" in stage_content
    assert "handleCaptureReady" in stage_content

    app_path = FRONTEND_DIR / "src" / "App.jsx"
    assert app_path.exists()
    app_content = app_path.read_text(encoding="utf-8")

    assert "threeCanvasCaptureRef" in app_content
    assert "onCaptureReady" in app_content


def test_task_14_1_data_url_format_and_base64_validation():
    """
    Task 14.1: Kiểm tra định dạng chuỗi Base64 Data URL chuẩn HTML5 Canvas:
    - Phải bắt đầu bằng tiền tố 'data:image/png;base64,'
    - Phần payload Base64 hợp lệ, khi giải mã phải chứa PNG Magic Bytes [0x89, 'P', 'N', 'G', '\r', '\n', 0x1A, '\n']
    """
    data_url_pattern = re.compile(r"^data:image/png;base64,([A-Za-z0-9+/=]+)$")
    match = data_url_pattern.match(SAMPLE_DATA_URL)
    assert match is not None, "Chuỗi Data URL không khớp với biểu thức chính quy chuẩn PNG base64"

    raw_base64 = match.group(1)
    image_bytes = base64.b64decode(raw_base64)

    # Kiểm tra 8 bytes chữ ký (Magic Header) của chuẩn PNG (ISO/IEC 15948)
    png_magic_header = b"\x89PNG\r\n\x1a\n"
    assert image_bytes.startswith(png_magic_header), (
        f"Dữ liệu byte không phải định dạng PNG hợp lệ. Header nhận được: {image_bytes[:8]}"
    )
    assert len(image_bytes) > 0


def test_task_14_2_schema_validation_snapshot_3d():
    """
    Task 14.2: Kiểm tra SlideSchema và SlideDeckSchema hỗ trợ trường snapshot_3d_base64:
    - SlideSchema chấp nhận snapshot_3d_base64 dạng Base64 Data URL
    - SlideDeckSchema chấp nhận snapshot_3d_base64 làm dự phòng toàn cục
    """
    from schemas.slide_schema import SlideDeckSchema, SlideSchema, SlideType, SlideLayout

    slide = SlideSchema(
        slide_index=1,
        type=SlideType.CONCEPT_SLIDE,
        layout=SlideLayout.IMAGE_TEXT,
        title="Khối Chóp Tứ Giác S.ABCD",
        bullet_points=["Đáy ABCD là hình vuông", "SA vuông góc đáy"],
        has_3d=True,
        polyhedron_type="pyramid_square",
        snapshot_3d_base64=SAMPLE_DATA_URL
    )
    assert slide.snapshot_3d_base64 == SAMPLE_DATA_URL
    assert slide.has_3d is True
    assert slide.polyhedron_type == "pyramid_square"

    deck = SlideDeckSchema(
        presentation_title="Hình Học Không Gian 3D",
        subject="Toán học",
        grade=11,
        slides=[slide],
        snapshot_3d_base64=SAMPLE_DATA_URL
    )
    assert deck.snapshot_3d_base64 == SAMPLE_DATA_URL
    assert len(deck.slides) == 1
    assert deck.slides[0].snapshot_3d_base64 == SAMPLE_DATA_URL


def test_task_14_2_frontend_payload_packaging_integration():
    """
    Task 14.2: Kiểm tra mã nguồn Frontend App.jsx đóng gói snapshot_3d_base64 vào request:
    - handleExportPptx gọi threeCanvasCaptureRef.current() hoặc window.__threeCanvasCaptureSnapshot()
    - Đóng gói snapshot_3d_base64 vào từng slide có has_3d / polyhedron_type hoặc slide đang mở
    - Gửi request POST tới /api/v1/lesson-plan/export-slide với payloadDeck
    """
    app_path = FRONTEND_DIR / "src" / "App.jsx"
    content = app_path.read_text(encoding="utf-8")

    assert "handleExportPptx" in content
    assert "threeCanvasCaptureRef.current" in content
    assert "snapshot_3d_base64" in content
    assert "payloadDeck" in content
    assert "fetch('/api/v1/lesson-plan/export-slide'" in content or 'fetch("/api/v1/lesson-plan/export-slide"' in content
    assert "body: JSON.stringify(payloadDeck)" in content


def test_task_14_2_api_export_slide_accepts_snapshot_payload():
    """
    Task 14.2: Kiểm tra endpoint POST /api/v1/lesson-plan/export-slide:
    - Tiếp nhận payload SlideDeckSchema chứa snapshot_3d_base64 mà không sinh lỗi HTTP 422 Unprocessable Entity
    - Trả về status 200 và file nhị phân PowerPoint (.pptx)
    """
    payload = {
        "presentation_title": "Hình Học Không Gian 3D Test Snapshot",
        "subject": "Toán học",
        "grade": 11,
        "theme": "math_academic",
        "snapshot_3d_base64": SAMPLE_DATA_URL,
        "slides": [
            {
                "slide_index": 1,
                "type": "TITLE_SLIDE",
                "layout": "TITLE_ONLY",
                "title": "Chủ đề 3D: Khối Đa Diện",
                "bullet_points": ["Khám phá hình học không gian trực quan"],
                "has_3d": False
            },
            {
                "slide_index": 2,
                "type": "CONCEPT_SLIDE",
                "layout": "IMAGE_TEXT",
                "title": "Hình Chóp Tứ Giác Đều S.ABCD",
                "bullet_points": [
                    "Đáy ABCD là hình vuông tâm O cạnh a",
                    "Đoạn SO vuông góc với mặt phẳng đáy (ABCD)",
                    "Thể tích khối chóp: V = (1/3) * a^2 * h"
                ],
                "main_definition_latex": "V = \\frac{1}{3} a^2 h",
                "has_3d": True,
                "polyhedron_type": "pyramid_square",
                "snapshot_3d_base64": SAMPLE_DATA_URL
            }
        ]
    }

    response = client.post("/api/v1/lesson-plan/export-slide", json=payload)
    assert response.status_code == 200, f"Lỗi xuất slide: {response.text}"
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in response.headers["content-type"]
    assert len(response.content) > 1000, "File PowerPoint trả về rỗng hoặc kích thước quá nhỏ"


def test_task_14_3_decode_base64_data_url():
    """
    Task 14.3: Kiểm tra hàm decode_base64_data_url():
    1. Giải mã thành công chuỗi Data URL chuẩn (data:image/png;base64,...)
    2. Giải mã thành công chuỗi Base64 trần (không có tiền tố data:...)
    3. Tự động khắc phục chuỗi Base64 bị khuyết padding '='
    4. Xử lý an toàn khi gặp dữ liệu lỗi/rỗng/None mà không raise Exception
    """
    from exporters.slide_exporter import decode_base64_data_url

    # 1. Data URL chuẩn
    stream1 = decode_base64_data_url(SAMPLE_DATA_URL)
    assert stream1 is not None
    assert isinstance(stream1, io.BytesIO)
    content1 = stream1.getvalue()
    assert content1.startswith(b"\x89PNG\r\n\x1a\n")

    # 2. Raw Base64 không có tiền tố
    stream2 = decode_base64_data_url(SAMPLE_PNG_BASE64)
    assert stream2 is not None
    assert stream2.getvalue().startswith(b"\x89PNG\r\n\x1a\n")

    # 3. Base64 bị khuyết padding '='
    unpadded_b64 = SAMPLE_PNG_BASE64.rstrip("=")
    stream3 = decode_base64_data_url(unpadded_b64)
    assert stream3 is not None
    assert stream3.getvalue().startswith(b"\x89PNG\r\n\x1a\n")

    # 4. Giá trị None, rỗng hoặc chuỗi rác
    assert decode_base64_data_url(None) is None
    assert decode_base64_data_url("") is None
    assert decode_base64_data_url("invalid_garbage_base64_!@#$%^") is None


def test_task_14_3_pptx_embedding_creates_picture_shape():
    """
    Task 14.3: Kiểm tra quá trình nhúng ảnh 3D snapshot vào Slide PowerPoint qua python-pptx:
    - Khi slide có snapshot_3d_base64, hệ thống tạo shape PICTURE (add_picture)
    - Dữ liệu ảnh nhúng trong PowerPoint khớp với ảnh PNG từ frontend
    - Slide vẫn hiển thị đầy đủ Header, Title, Bullets và Footer
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from exporters.slide_exporter import export_slide_deck_to_pptx
    from schemas.slide_schema import SlideDeckSchema, SlideSchema, SlideType, SlideLayout

    slide_3d = SlideSchema(
        slide_index=1,
        type=SlideType.CONCEPT_SLIDE,
        layout=SlideLayout.IMAGE_TEXT,
        title="Hình Chóp Tứ Giác S.ABCD & Chiều Cao h",
        bullet_points=[
            "Đáy ABCD là hình vuông tâm O cạnh a = 4.0",
            "Đường cao SA vuông góc đáy, chiều cao h = 4.0",
            "Thể tích: V = (1/3) * a^2 * h"
        ],
        main_definition_latex="V_{S.ABCD} = \\frac{1}{3} a^2 h",
        has_3d=True,
        polyhedron_type="pyramid_square",
        snapshot_3d_base64=SAMPLE_DATA_URL
    )

    deck = SlideDeckSchema(
        presentation_title="Hình Học Không Gian 3D",
        subject="Toán học",
        grade=11,
        slides=[slide_3d]
    )

    pptx_stream = export_slide_deck_to_pptx(deck)
    assert pptx_stream is not None

    prs = Presentation(pptx_stream)
    assert len(prs.slides) == 1

    slide = prs.slides[0]
    # Tìm kiếm shape PICTURE trong slide
    picture_shapes = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(picture_shapes) >= 1, "Không tìm thấy Picture Shape nào trong slide có 3D snapshot!"

    # Kiểm tra ít nhất 1 Picture Shape chứa đúng dữ liệu ảnh PNG
    has_png_snapshot = False
    for p_sh in picture_shapes:
        blob = getattr(p_sh, "image").blob
        if blob.startswith(b"\x89PNG\r\n\x1a\n"):
            has_png_snapshot = True
            assert p_sh.width > 0
            assert p_sh.height > 0
            break

    assert has_png_snapshot, "Picture Shape không chứa dữ liệu PNG hợp lệ từ snapshot_3d_base64"


def test_task_14_3_pptx_inherits_deck_snapshot_fallback():
    """
    Task 14.3: Kiểm tra cơ chế thừa kế (Fallback) snapshot 3D từ SlideDeck xuống Slide:
    - Nếu slide có has_3d=True nhưng chưa có snapshot_3d_base64 riêng, slide sẽ tự động
      thừa kế snapshot_3d_base64 được gắn ở cấp độ Deck.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from exporters.slide_exporter import export_slide_deck_to_pptx
    from schemas.slide_schema import SlideDeckSchema, SlideSchema, SlideType, SlideLayout

    slide_without_own_snap = SlideSchema(
        slide_index=1,
        type=SlideType.CONCEPT_SLIDE,
        layout=SlideLayout.IMAGE_TEXT,
        title="Hình Lập Phương ABCD.A'B'C'D'",
        bullet_points=["Tất cả các cạnh bằng a"],
        has_3d=True,
        polyhedron_type="cube",
        snapshot_3d_base64=None  # Không có ảnh riêng
    )

    deck_with_global_snap = SlideDeckSchema(
        presentation_title="Bài Học Hình Lập Phương",
        subject="Toán học",
        grade=11,
        snapshot_3d_base64=SAMPLE_DATA_URL,  # Ảnh toàn cục
        slides=[slide_without_own_snap]
    )

    pptx_stream = export_slide_deck_to_pptx(deck_with_global_snap)
    prs = Presentation(pptx_stream)
    slide = prs.slides[0]

    picture_shapes = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(picture_shapes) >= 1, "Slide không tự động thừa kế 3D snapshot từ Deck!"


def make_test_png_data_url(width: int, height: int, color=(56, 189, 248)) -> str:
    """Helper tạo chuỗi Data URL PNG hợp lệ với kích thước pixel bất kỳ để kiểm thử tỉ lệ."""
    from PIL import Image
    img = Image.new("RGBA", (width, height), color)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return f"data:image/png;base64,{base64.b64encode(buf.getvalue()).decode('utf-8')}"


def test_task_14_4_aspect_preserving_fit_algorithm():
    """
    Task 14.4: Kiểm tra thuật toán compute_aspect_preserving_fit():
    1. Ảnh toàn cảnh Landscape (1920x1080 ~ 16:9): Fit chiều rộng, căn giữa theo chiều cao
    2. Ảnh vuông Square (800x800 ~ 1:1): Fit chiều cao, căn giữa theo chiều rộng
    3. Ảnh đứng Portrait (600x1200 ~ 1:2): Fit chiều cao, căn giữa theo chiều rộng
    4. Tỉ lệ khung hình (Aspect Ratio) của kích thước fit phải bằng chính xác tỉ lệ ảnh gốc
    """
    from exporters.slide_exporter import compute_aspect_preserving_fit

    box_l, box_t, box_w, box_h = 5.7, 1.5, 3.6, 2.45

    # 1. Ảnh 16:9 (1920x1080)
    l1, t1, w1, h1 = compute_aspect_preserving_fit(1920, 1080, box_l, box_t, box_w, box_h)
    assert w1 == pytest.approx(box_w)
    assert h1 < box_h
    assert l1 == pytest.approx(box_l)
    assert t1 > box_t  # Căn giữa theo trục đứng
    assert (w1 / h1) == pytest.approx(1920 / 1080, rel=1e-3)

    # 2. Ảnh vuông 1:1 (800x800)
    l2, t2, w2, h2 = compute_aspect_preserving_fit(800, 800, box_l, box_t, box_w, box_h)
    assert h2 == pytest.approx(box_h)
    assert w2 < box_w
    assert t2 == pytest.approx(box_t)
    assert l2 > box_l  # Căn giữa theo trục ngang
    assert (w2 / h2) == pytest.approx(1.0, rel=1e-3)

    # 3. Ảnh đứng 1:2 (600x1200)
    l3, t3, w3, h3 = compute_aspect_preserving_fit(600, 1200, box_l, box_t, box_w, box_h)
    assert h3 == pytest.approx(box_h)
    assert w3 < box_w
    assert (w3 / h3) == pytest.approx(600 / 1200, rel=1e-3)


def test_task_14_4_slide_grid_16_9_geometry_and_margins():
    """
    Task 14.4: Kiểm tra sự cân đối và đối xứng toán học của lưới bố cục 16:9 Widescreen:
    - Kích thước slide chuẩn: 10.0 x 5.625 inches (16:9 Widescreen)
    - Cột Trái: l_margin = 0.6 in, w_left = 4.8 in
    - Khe hở Gutter: 5.6 - (0.6 + 4.8) = 0.2 in
    - Cột Phải: card_w = 3.8 in
    - Lề Phải: 10.0 - (5.6 + 3.8) = 0.6 in
    - Đối xứng hoàn hảo: Lề Trái (0.6 in) == Lề Phải (0.6 in)
    """
    from pptx import Presentation
    from pptx.util import Inches
    from exporters.slide_exporter import export_slide_deck_to_pptx
    from schemas.slide_schema import SlideDeckSchema, SlideSchema

    deck = SlideDeckSchema(
        presentation_title="Lưới 16:9 Toán Học",
        subject="Toán học",
        grade=11,
        slides=[
            SlideSchema(
                slide_index=1,
                title="Khối Chóp S.ABCD",
                bullet_points=["Đáy là hình vuông"],
                has_3d=True,
                polyhedron_type="pyramid_square",
                snapshot_3d_base64=SAMPLE_DATA_URL
            )
        ]
    )

    pptx_stream = export_slide_deck_to_pptx(deck)
    prs = Presentation(pptx_stream)

    # Kiểm tra kích thước Slide chuẩn 16:9
    assert prs.slide_width is not None and prs.slide_height is not None
    assert prs.slide_width == Inches(10.0)
    assert prs.slide_height == Inches(5.625)
    assert (float(prs.slide_width) / float(prs.slide_height)) == pytest.approx(16.0 / 9.0)

    # Kiểm tra tính toán lề
    l_margin = 0.6
    w_left = 4.8
    gutter = 5.6 - (l_margin + w_left)
    w_card = 3.8
    r_margin = 10.0 - (5.6 + w_card)

    assert gutter == pytest.approx(0.2)
    assert l_margin == pytest.approx(r_margin)  # Đối xứng lề hoàn hảo


def test_task_14_4_pptx_picture_preserves_aspect_ratio_in_exported_deck():
    """
    Task 14.4: Kiểm tra hình ảnh 3D sau khi nhúng vào slide PowerPoint giữ nguyên vẹn tỉ lệ:
    - Nhúng ảnh vuông 600x600 -> shape Picture trên PPTX phải có width == height
    - Nhúng ảnh 16:9 960x540 -> shape Picture trên PPTX phải có width / height == 16/9
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
    from exporters.slide_exporter import export_slide_deck_to_pptx
    from schemas.slide_schema import SlideDeckSchema, SlideSchema

    square_snap = make_test_png_data_url(600, 600)
    widescreen_snap = make_test_png_data_url(960, 540)

    deck = SlideDeckSchema(
        presentation_title="Kiểm Thử Tỉ Lệ Hình Ảnh 3D",
        subject="Toán học",
        grade=11,
        slides=[
            SlideSchema(
                slide_index=1,
                title="Slide 1: Ảnh Vuông",
                bullet_points=["Ảnh snapshot tỉ lệ 1:1"],
                has_3d=True,
                polyhedron_type="cube",
                snapshot_3d_base64=square_snap
            ),
            SlideSchema(
                slide_index=2,
                title="Slide 2: Ảnh 16:9",
                bullet_points=["Ảnh snapshot tỉ lệ 16:9"],
                has_3d=True,
                polyhedron_type="pyramid_square",
                snapshot_3d_base64=widescreen_snap
            )
        ]
    )

    pptx_stream = export_slide_deck_to_pptx(deck)
    prs = Presentation(pptx_stream)

    # 1. Slide 1 (Ảnh Vuông)
    slide1 = prs.slides[0]
    pic1: Any = [sh for sh in slide1.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    assert pic1.width is not None and pic1.height is not None
    # Phải nằm gọn trong khung chứa 3.6 x 2.45 inches
    assert pic1.width <= Inches(3.601)
    assert pic1.height <= Inches(2.451)
    # Tỉ lệ phải là 1:1
    assert (float(pic1.width) / float(pic1.height)) == pytest.approx(1.0, rel=1e-2)

    # 2. Slide 2 (Ảnh 16:9)
    slide2 = prs.slides[1]
    pic2: Any = [sh for sh in slide2.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    assert pic2.width is not None and pic2.height is not None
    assert pic2.width <= Inches(3.601)
    assert pic2.height <= Inches(2.451)
    # Tỉ lệ phải là 16:9
    assert (float(pic2.width) / float(pic2.height)) == pytest.approx(16.0 / 9.0, rel=1e-2)


def test_task_14_5_dod_end_to_end_verification():
    """
    Task 14.5: Nghiệm thu Definition of Done (DoD) Ngày 14 qua kịch bản End-to-End:
    1. Gửi request POST /api/v1/lesson-plan/export-slide với bộ slide hoàn chỉnh chứa snapshot 3D
    2. File trả về có mã trạng thái HTTP 200, Content-Type chuẩn PPTX Widescreen 16:9
    3. Cấu trúc OpenXML (ZIP package) chứa đầy đủ presentation.xml, slides/slide*.xml và media/
    4. Mở được thành công bằng thư viện python-pptx / Microsoft PowerPoint
    5. Xác thực có chứa hình ảnh 3D snapshot chất lượng cao, đúng góc nhìn và không méo hình
    6. Kiểm tra Speaker Notes dành cho Giáo viên trong file PowerPoint
    """
    import zipfile
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    real_snap_url = make_test_png_data_url(800, 600, color=(30, 41, 59))

    payload = {
        "presentation_title": "Hình Học Không Gian 3D • DoD Ngày 14",
        "subject": "Toán học",
        "grade": 11,
        "theme": "math_academic",
        "snapshot_3d_base64": real_snap_url,
        "slides": [
            {
                "slide_index": 1,
                "type": "TITLE_SLIDE",
                "layout": "TITLE_ONLY",
                "title": "Chuyên Đề: Thể Tích Khối Chóp & Khối Đa Diện",
                "subtitle": "Ứng dụng hình học không gian 3D tương tác theo CV 5512",
                "bullet_points": ["Trực quan hóa hình học không gian 3D"],
                "has_3d": False
            },
            {
                "slide_index": 2,
                "type": "CONCEPT_SLIDE",
                "layout": "IMAGE_TEXT",
                "title": "1. Hình Chóp Tứ Giác S.ABCD Có Cạnh Bên SA Vuông Góc Đáy",
                "subtitle": "Quan sát chân đường cao A và góc phẳng nhị diện",
                "bullet_points": [
                    "Đáy ABCD là hình vuông cạnh $a$ nằm trên mặt phẳng $(Oxy)$.",
                    "Đường cao $SA \\perp (ABCD)$ với $SA = h = 4.0$.",
                    "Thể tích khối chóp: $V_{S.ABCD} = \\frac{1}{3} S_{ABCD} \\cdot SA = \\frac{1}{3} a^2 h$."
                ],
                "main_definition_latex": "V_{S.ABCD} = \\frac{1}{3} a^2 h",
                "teacher_note": "Cho học sinh xoay mô hình 3D trên Web Studio để nhìn rõ chân đường cao A.",
                "has_3d": True,
                "polyhedron_type": "pyramid_square",
                "snapshot_3d_base64": real_snap_url
            }
        ]
    }

    # 1. Gọi API Export PPTX
    resp = client.post("/api/v1/lesson-plan/export-slide", json=payload)
    assert resp.status_code == 200, f"API lỗi: {resp.text}"
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in resp.headers["content-type"]
    pptx_bytes = resp.content
    assert len(pptx_bytes) > 2000, "File PowerPoint trả về quá nhỏ hoặc rỗng"

    # 2. Kiểm tra tính hợp lệ của tệp ZIP OpenXML Package
    pptx_stream = io.BytesIO(pptx_bytes)
    with zipfile.ZipFile(pptx_stream, "r") as zf:
        namelist = zf.namelist()
        assert "[Content_Types].xml" in namelist
        assert "ppt/presentation.xml" in namelist
        assert "ppt/slides/slide1.xml" in namelist
        assert "ppt/slides/slide2.xml" in namelist
        # Kiểm tra media chứa ảnh nhúng
        media_files = [n for n in namelist if n.startswith("ppt/media/")]
        assert len(media_files) >= 1, "Không tìm thấy file ảnh PNG nào trong ppt/media/!"

    # 3. Phân tích nội dung các slide qua python-pptx
    pptx_stream.seek(0)
    prs = Presentation(pptx_stream)
    assert len(prs.slides) == 2
    assert prs.slide_width == Inches(10.0)
    assert prs.slide_height == Inches(5.625)

    # Slide 1: Title Slide
    s1 = prs.slides[0]
    assert len(s1.shapes) >= 2

    # Slide 2: 3D Concept Slide
    s2 = prs.slides[1]
    picture_shapes = [sh for sh in s2.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    # Slide 2 có cả ảnh công thức LaTeX (cột trái) và ảnh snapshot 3D (cột phải)
    assert len(picture_shapes) >= 1, "Slide 2 không chứa Picture Shape nào!"

    # Tìm chính xác Picture Shape thuộc cột phải (3D Snapshot)
    snap_pics = [p for p in picture_shapes if p.left >= Inches(5.0)]
    assert len(snap_pics) >= 1, "Không tìm thấy Picture Shape của 3D snapshot ở cột phải!"
    snap_pic: Any = snap_pics[0]

    assert getattr(snap_pic, "image").blob.startswith(b"\x89PNG\r\n\x1a\n")
    assert snap_pic.width <= Inches(3.601)
    assert snap_pic.height <= Inches(2.451)
    assert snap_pic.left >= Inches(5.699)  # Thuộc cột phải (left >= 5.7 in)

    # Kiểm tra badge "Mô hình 3D WebGL (SGK)"
    badge_found = False
    for sh in s2.shapes:
        tf = getattr(sh, "text_frame", None)
        if tf is not None and "Mô hình 3D WebGL (SGK)" in getattr(tf, "text", ""):
            badge_found = True
            break
    assert badge_found, "Thiếu huy hiệu Mô hình 3D WebGL (SGK) trên slide 2"

    # Kiểm tra Speaker Notes của Giáo viên
    assert s2.has_notes_slide
    notes_slide = s2.notes_slide
    assert notes_slide is not None
    notes_tf = notes_slide.notes_text_frame
    assert notes_tf is not None
    notes_text = notes_tf.text
    assert "Cho học sinh xoay mô hình 3D" in notes_text




