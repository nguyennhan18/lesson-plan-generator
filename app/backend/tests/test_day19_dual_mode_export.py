# -*- coding: utf-8 -*-
"""
Test Suite Ngày 19: Đồng Bộ Xuất File Word (docx_exporter.py) & Hoàn Thiện Luồng Dual-Mode
Task 19.1: Nâng cấp docx_exporter.py & Hiện thực Thuật toán Đồng bộ 100% Nội dung Giáo án Word (CV 5512) với Slide Bài giảng
"""

import os
import io
import sys
import time
import zipfile
from pathlib import Path
import pytest
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from fastapi.testclient import TestClient

ROOT_DIR = Path(__file__).resolve().parents[3]
BACKEND_SRC = ROOT_DIR / "app" / "backend" / "src"

if str(BACKEND_SRC) not in sys.path:
    sys.path.insert(0, str(BACKEND_SRC))

from schemas.schemas import LessonPlanSchema, LessonActivity, ExecutionSteps
from schemas.slide_schema import SlideDeckSchema
from exporters.docx_exporter import (
    export_lesson_plan_to_docx,
    export_lesson_plan_to_docx_file,
    synchronize_lesson_and_slide_deck,
    clean_xml_string,
    repair_latex_code,
    convert_latex_to_omml,
    add_text_with_latex,
)
from exporters.dual_package_exporter import (
    export_dual_package,
    remove_vietnamese_diacritics,
    sanitize_filename,
    DualPackageResult,
)
from utils.temp_file_cleaner import (
    TempFileCleaner,
    CleaningReport,
    TempStatsReport,
    clean_expired_temp_files,
    cleanup_memory_and_figures,
    get_temp_storage_stats,
)
from main import app

client = TestClient(app)


@pytest.fixture
def sample_lesson_plan():
    """Tạo đối tượng Kế hoạch bài dạy chuẩn CV 5512"""
    return LessonPlanSchema(
        lesson_title="Thể Tích Khối Lăng Trụ Đứng",
        subject="Toán học",
        grade=11,
        knowledge_goals=[
            r"Nắm vững định nghĩa và tính chất của hình lăng trụ đứng tam giác và tứ giác.",
            r"Thuộc và giải thích được ý nghĩa các đại lượng trong công thức tính thể tích \(V = B \cdot h\).",
            r"Biết cách tính diện tích đáy tam giác đều cạnh a: \(S = \frac{a^2\sqrt{3}}{4}\)."
        ],
        competency_goals=[
            "Năng lực tư duy và lập luận toán học: Chứng minh các cạnh bên vuông góc với đáy.",
            "Năng lực giải quyết vấn đề toán học: Tính toán thể tích chính xác không nhầm lẫn với hình chóp."
        ],
        character_goals=[
            "Chăm chỉ: Tích cực hoàn thành các bài tập rèn luyện cá nhân và theo nhóm.",
            "Trung thực: Nghiêm túc trong quá trình làm bài kiểm tra trắc nghiệm."
        ],
        teaching_equipment=[
            "Sách giáo khoa Toán 11 - Kết nối tri thức.",
            "Máy tính, máy chiếu và phần mềm Geogebra."
        ],
        activities=[
            LessonActivity(
                activity_number=1,
                time_minutes=5,
                goal="Nhắc lại khái niệm hình chóp và đặt vấn đề thể tích khối lăng trụ.",
                expected_product="Câu trả lời của học sinh về mối quan hệ giữa chiều cao và cạnh đáy.",
                execution=ExecutionSteps(
                    step_1_assign="GV chiếu hình ảnh lăng trụ trong thực tế và đặt câu hỏi gợi mở.",
                    step_2_excute="HS quan sát hình ảnh và trao đổi theo bàn.",
                    step_3_report="Đại diện học sinh trả lời, các học sinh khác nhận xét.",
                    step_4_conclusion="GV nhận xét và dẫn dắt vào bài mới."
                )
            ),
            LessonActivity(
                activity_number=2,
                time_minutes=20,
                goal="Hình thành công thức tính thể tích khối lăng trụ.",
                expected_product=r"Công thức tính thể tích \(V = B \cdot h\) với B là diện tích đáy, h là chiều cao.",
                execution=ExecutionSteps(
                    step_1_assign="GV hướng dẫn học sinh đọc SGK và quan sát mô hình trực quan.",
                    step_2_excute="HS thực hiện khám phá công thức thể tích qua hình lăng trụ đứng.",
                    step_3_report="Đại diện nhóm báo cáo kết luận.",
                    step_4_conclusion=r"GV chốt lại công thức \(V = B \cdot h\) và lưu ý đơn vị đo."
                )
            ),
            LessonActivity(
                activity_number=3,
                time_minutes=15,
                goal="Luyện tập củng cố công thức tính thể tích khối lăng trụ đứng.",
                expected_product="Bài giải ví dụ mẫu và đáp án câu hỏi trắc nghiệm của học sinh.",
                execution=ExecutionSteps(
                    step_1_assign="GV giao bài tập ví dụ và phiếu trắc nghiệm nhanh cho học sinh.",
                    step_2_excute="HS làm bài tập cá nhân vào vở.",
                    step_3_report="HS lên bảng trình bày lời giải và chọn đáp án trắc nghiệm.",
                    step_4_conclusion="GV nhận xét, sửa lỗi sai và phân tích các bẫy thường gặp."
                )
            ),
            LessonActivity(
                activity_number=4,
                time_minutes=5,
                goal="Vận dụng tính thể tích một khối bê tông lăng trụ trong thực tiễn.",
                expected_product="Lời giải bài toán tính lượng bê tông đổ cột nhà hình lăng trụ.",
                execution=ExecutionSteps(
                    step_1_assign="GV giao nhiệm vụ về nhà cho học sinh.",
                    step_2_excute="HS ghi lại đề bài vào vở ghi.",
                    step_3_report="Báo cáo kết quả vào đầu tiết học tiếp theo.",
                    step_4_conclusion="GV đánh giá tiết học và dặn dò ôn bài."
                )
            )
        ]
    )


@pytest.fixture
def sample_slide_deck_day19():
    """Tạo bộ Slide trình chiếu đồng bộ để kiểm thử tích hợp Word ↔ PPTX"""
    return {
        "presentation_title": "Thể Tích Khối Lăng Trụ Đứng & Ứng Dụng",
        "subject": "Toán học",
        "grade": 11,
        "slides": [
            {
                "slide_index": 1,
                "type": "TITLE_SLIDE",
                "title": "Thể Tích Khối Lăng Trụ Đứng",
                "subtitle": "Chương trình Hình học Không gian Lớp 11",
            },
            {
                "slide_index": 2,
                "type": "CONCEPT_SLIDE",
                "title": "Định Nghĩa & Công Thức Tính Thể Tích",
                "bullet_points": [
                    "Thể tích khối lăng trụ bằng diện tích đáy nhân với chiều cao.",
                    "Lăng trụ đứng có chiều cao h chính bằng độ dài cạnh bên."
                ],
                "main_definition_latex": r"V = B \cdot h = S_{\Delta ABC} \cdot AA'",
                "has_3d": True,
                "polyhedron_type": "prism_triangle",
            },
            {
                "slide_index": 3,
                "type": "EXERCISE_SLIDE",
                "title": "Ví Dụ Áp Dụng: Lăng Trụ Tam Giác Đều",
                "worked_solution_steps": [
                    r"Bước 1: Tính diện tích đáy tam giác đều cạnh a: $S = \frac{a^2\sqrt{3}}{4}$",
                    r"Bước 2: Xác định chiều cao lăng trụ đứng: $h = AA' = 2a$",
                    r"Bước 3: Thể tích khối lăng trụ: $V = S \cdot h = \frac{a^3\sqrt{3}}{2}$"
                ],
                "teacher_note": "Nhắc nhở học sinh trình bày đầy đủ các bước tự luận."
            },
            {
                "slide_index": 4,
                "type": "EXERCISE_SLIDE",
                "title": "Câu Hỏi Trắc Nghiệm Củng Cố",
                "layout": "QUIZ_OPTION",
                "activity_type": "Luyện tập trắc nghiệm",
                "question": r"Cho hình lăng trụ đứng có diện tích đáy $B = 15\,\text{cm}^2$ và chiều cao $h = 4\,\text{cm}$. Thể tích khối lăng trụ là:",
                "quiz_options": [
                    {"label": "A", "text": r"$20\,\text{cm}^3$", "is_correct": False, "distractor_rationale": "Nhầm công thức thể tích khối chóp có hệ số 1/3."},
                    {"label": "B", "text": r"$60\,\text{cm}^3$", "is_correct": True, "distractor_rationale": ""},
                    {"label": "C", "text": r"$30\,\text{cm}^3$", "is_correct": False, "distractor_rationale": "Chia đôi diện tích do nhầm công thức diện tích tam giác."},
                    {"label": "D", "text": r"$120\,\text{cm}^3$", "is_correct": False, "distractor_rationale": "Nhân đôi chiều cao h."}
                ],
                "teacher_note": "Cho học sinh 60 giây suy nghĩ trước khi chọn đáp án."
            }
        ]
    }


def test_task_19_1_export_pure_lesson_plan_to_docx(sample_lesson_plan):
    """
    Task 19.1: Kiểm tra xuất file Word từ đối tượng LessonPlanSchema độc lập:
    - Đúng quy chuẩn lề hành chính Việt Nam (Nghị định 30/2020/NĐ-CP): Trái 1.18", Phải 0.59", Trên/Dưới 0.79"
    - Cấu trúc đầy đủ: Tiêu đề, Phần I, Phần II, Phần III (Bảng 4 bước cho 4 hoạt động), Khung ký duyệt
    - Định dạng font chữ mặc định: Times New Roman, size 13 pt
    """
    docx_stream = export_lesson_plan_to_docx(sample_lesson_plan)
    assert isinstance(docx_stream, io.BytesIO)
    assert docx_stream.getbuffer().nbytes > 10000, "Tệp Word sinh ra quá nhỏ"

    doc = Document(docx_stream)

    # 1. Kiểm tra lề trang
    sec = doc.sections[0]
    assert sec.left_margin is not None and sec.right_margin is not None
    assert sec.top_margin is not None and sec.bottom_margin is not None
    assert round(sec.left_margin.inches, 2) == 1.18
    assert round(sec.right_margin.inches, 2) == 0.59
    assert round(sec.top_margin.inches, 2) == 0.79
    assert round(sec.bottom_margin.inches, 2) == 0.79

    # 2. Kiểm tra các tiêu đề chính
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "KẾ HOẠCH BÀI DẠY (GIÁO ÁN)" in full_text
    assert "MÔN: TOÁN HỌC - LỚP 11" in full_text
    assert "BÀI HỌC: THỂ TÍCH KHỐI LĂNG TRỤ ĐỨNG" in full_text
    assert "I. MỤC TIÊU BÀI HỌC" in full_text
    assert "II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU" in full_text
    assert "III. TIẾN TRÌNH DẠY HỌC" in full_text
    assert "TỔ TRƯỞNG CHUYÊN MÔN DUYỆT" in "\n".join([c.text for row in doc.tables[-1].rows for c in row.cells])

    # 3. Kiểm tra bảng hoạt động (mỗi hoạt động có 1 bảng 5 hàng x 2 cột)
    assert len(doc.tables) >= 5  # 4 bảng hoạt động + 1 bảng ký duyệt
    act_table = doc.tables[0]
    assert len(act_table.rows) == 5
    assert len(act_table.columns) == 2
    assert "Các bước tổ chức" in act_table.rows[0].cells[0].text
    assert "Bước 1: Chuyển giao nhiệm vụ" in act_table.rows[1].cells[0].text
    assert "Bước 4: Kết luận & nhận định" in act_table.rows[4].cells[0].text


def test_task_19_1_omml_native_equation_and_repair():
    """
    Task 19.1: Kiểm tra chuyển đổi công thức Toán học sang XML OMML chuẩn Microsoft Word:
    - Sửa cú pháp: rac{a}{b} -> \frac{a}{b}, ext{cm} -> \text{cm}
    - Tạo thành công đối tượng parse_xml chứa namespace xmlns:m
    - Nhúng vào Paragraph dưới dạng native math element
    """
    # 1. Test auto-repair
    assert repair_latex_code(r"rac{1}{2}") == r"\frac{1}{2}"
    assert repair_latex_code(r"20ext{ cm}^3") == r"20\text{ cm}^3"

    # 2. Test convert_latex_to_omml
    omml1 = convert_latex_to_omml(r"\frac{a^2\sqrt{3}}{4}")
    assert omml1 is not None
    assert "oMath" in omml1.tag or "oMath" in str(omml1)

    omml2 = convert_latex_to_omml(r"V = B \cdot h")
    assert omml2 is not None

    # 3. Test add_text_with_latex nhúng vào Paragraph
    doc = Document()
    p = doc.add_paragraph()
    add_text_with_latex(p, r"Công thức tính thể tích là $V = B \cdot h$ và diện tích đáy $S = \frac{a^2\sqrt{3}}{4}$.")
    # Kiểm tra trong XML của paragraph có chứa thẻ OMML m:oMath
    p_xml = p._p.xml
    assert "oMath" in p_xml
    assert "Times New Roman" in p_xml


def test_task_19_1_bi_directional_synchronization(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.1: Kiểm tra giải thuật đồng bộ dữ liệu song hướng (synchronize_lesson_and_slide_deck):
    - Tự động bổ sung thiết bị 3D WebGL (Three.js) vào Phần II
    - Tự động bổ sung công thức cốt lõi V = B * h vào Hoạt động 2
    - Tự động bổ sung ghi chú số lượng câu hỏi trắc nghiệm và bài tập vào Hoạt động 3
    """
    plan_synced, deck_info = synchronize_lesson_and_slide_deck(sample_lesson_plan, sample_slide_deck_day19)

    # 1. Kiểm tra deck_info bóc tách đúng thành phần
    assert len(deck_info["quiz_items"]) == 1
    assert len(deck_info["worked_examples"]) == 1
    assert len(deck_info["concept_formulas"]) == 1
    assert len(deck_info["visual_models"]) == 1
    assert deck_info["visual_models"][0]["polyhedron_type"] == "prism_triangle"

    # 2. Kiểm tra thiết bị dạy học được bổ sung
    eq_str = "\n".join(plan_synced.teaching_equipment)
    assert "Mô hình Hình học không gian 3D tương tác WebGL" in eq_str
    assert "Bài giảng điện tử tương tác chuẩn Web Studio" in eq_str

    # 3. Kiểm tra Hoạt động 2 chứa công thức toán
    act2_prod = plan_synced.activities[1].expected_product
    assert r"V = B \cdot h = S_{\Delta ABC} \cdot AA'" in act2_prod

    # 4. Kiểm tra Hoạt động 3 chứa ghi chú trắc nghiệm và ví dụ tự luận
    act3_goal = plan_synced.activities[2].goal
    assert "1 câu hỏi trắc nghiệm rèn luyện" in act3_goal
    assert "1 bài tập tự luận mẫu" in act3_goal


def test_task_19_1_docx_embedded_quiz_and_worked_solutions(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.1: Kiểm tra việc nhúng trọn vẹn Câu hỏi Trắc nghiệm & Lời giải Tự luận vào file Word:
    - Bảng câu hỏi trắc nghiệm đồng bộ Slide hiển thị ở Hoạt động 3
    - Đầy đủ 4 phương án [A], [B], [C], [D]
    - Nêu rõ [Đáp án đúng chuẩn SGK]: Phương án B
    - Bóc tách phân tích bẫy sai lầm của các phương án nhiễu A, C, D
    - Hướng dẫn giải tự luận chi tiết từng bước: Bước 1, Bước 2, Bước 3
    """
    docx_stream = export_lesson_plan_to_docx(sample_lesson_plan, slide_deck=sample_slide_deck_day19)
    doc = Document(docx_stream)
    full_text = "\n".join([p.text for p in doc.paragraphs])

    # 1. Trắc nghiệm đồng bộ
    assert "HỆ THỐNG CÂU HỎI TRẮC NGHIỆM ĐÁNH GIÁ ĐỊNH DẠNG BGD&ĐT (ĐỒNG BỘ SLIDE):" in full_text
    assert "Cho hình lăng trụ đứng có diện tích đáy" in full_text
    assert "[A]." in full_text
    assert "[B]." in full_text
    assert "[C]." in full_text
    assert "[D]." in full_text
    assert "[Đáp án đúng chuẩn SGK]: Phương án B" in full_text
    assert "[Phân tích bẫy sai lầm học sinh dễ mắc phải]:" in full_text
    assert "Phương án A:" in full_text
    assert "Nhầm công thức thể tích khối chóp có hệ số 1/3" in full_text
    assert "Phương án C:" in full_text
    assert "Phương án D:" in full_text

    # 2. Lời giải mẫu tự luận đồng bộ
    assert "HƯỚNG DẪN GIẢI MẪU TỰ LUẬN TỪNG BƯỚC CHUẨN MỰC SGK:" in full_text
    assert "Ví dụ 1: Ví Dụ Áp Dụng: Lăng Trụ Tam Giác Đều" in full_text
    assert "Bước 1: Tính diện tích đáy tam giác đều" in full_text
    assert "Bước 2: Xác định chiều cao lăng trụ đứng" in full_text
    assert "Bước 3: Thể tích khối lăng trụ" in full_text


def test_task_19_1_xml_sanitization_in_docx():
    """
    Task 19.1: Kiểm tra lọc ký tự điều khiển cấm trong Word (clean_xml_string):
    - Đảm bảo các ký tự \x00-\x08, \x0B-\x0C, \x0E-\x1F bị loại bỏ triệt để
    - File Word sinh ra không bị corrupt cấu trúc XML OpenXML
    """
    dirty_text = "Giáo án Toán\x00\x01\x05\x0B\x0CLớp 11\x1E\x1F - GDPT 2018"
    cleaned = clean_xml_string(dirty_text)
    assert cleaned == "Giáo án ToánLớp 11 - GDPT 2018"

    plan = LessonPlanSchema(
        lesson_title=dirty_text,
        subject="Toán\x00 học",
        grade=11,
        knowledge_goals=["Mục tiêu\x05 1", "Mục tiêu\x0C 2"],
        activities=[]
    )
    docx_stream = export_lesson_plan_to_docx(plan)
    assert isinstance(docx_stream, io.BytesIO)
    doc = Document(docx_stream)
    title_text = doc.paragraphs[2].text
    for ch in title_text:
        code = ord(ch)
        assert not (0 <= code <= 8 or 11 <= code <= 12 or 14 <= code <= 31)


def test_task_19_1_export_docx_endpoint_api(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.1: Kiểm tra API POST /api/v1/lesson-plan/export-docx:
    - Nhận payload chứa cả plan và slide_deck đồng bộ
    - Trả về status 200, Content-Type chuẩn docx
    - Tệp tin nhị phân tải về mở được bằng python-docx và chứa đầy đủ nội dung đồng bộ
    """
    payload = {
        "plan": sample_lesson_plan.model_dump(),
        "slide_deck": sample_slide_deck_day19,
    }
    response = client.post("/api/v1/lesson-plan/export-docx", json=payload)
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers["content-type"]
    assert "attachment; filename" in response.headers["content-disposition"]

    docx_bytes = io.BytesIO(response.content)
    doc = Document(docx_bytes)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "KẾ HOẠCH BÀI DẠY (GIÁO ÁN)" in full_text
    assert "HỆ THỐNG CÂU HỎI TRẮC NGHIỆM ĐÁNH GIÁ ĐỊNH DẠNG BGD&ĐT" in full_text


def test_task_19_2_export_dual_package_zip_endpoint(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.2: Kiểm tra API POST /api/v1/lesson-plan/export-dual-package:
    - Nhận payload gồm cả plan và slide_deck
    - Trả về tệp tin ZIP chuẩn nén (media_type="application/zip")
    - Đóng gói đầy đủ 3 thành phần:
        1. Giao_An_5512_{title}.docx (hợp chuẩn python-docx, chứa nội dung đồng bộ)
        2. Slide_Bai_Giang_{title}.pptx (hợp chuẩn python-pptx, chứa đầy đủ các slide)
        3. Huong_Dan_Goi_Kep_5512.txt (hướng dẫn tích hợp sư phạm)
    """
    payload = {
        "plan": sample_lesson_plan.model_dump(),
        "slide_deck": sample_slide_deck_day19
    }
    response = client.post("/api/v1/lesson-plan/export-dual-package", json=payload)
    if response.status_code != 200:
        print("ERROR DUAL PACKAGE:", response.status_code, response.text)
    assert response.status_code == 200
    assert "application/zip" in response.headers["content-type"]
    assert "attachment; filename" in response.headers["content-disposition"]
    assert ".zip" in response.headers["content-disposition"]

    zip_bytes = io.BytesIO(response.content)
    with zipfile.ZipFile(zip_bytes, "r") as zf:
        namelist = zf.namelist()
        assert len(namelist) == 3

        docx_names = [n for n in namelist if n.endswith(".docx")]
        pptx_names = [n for n in namelist if n.endswith(".pptx")]
        txt_names = [n for n in namelist if n.endswith(".txt")]

        assert len(docx_names) == 1
        assert len(pptx_names) == 1
        assert len(txt_names) == 1

        # Kiểm tra file Word trong ZIP
        docx_data = zf.read(docx_names[0])
        doc = Document(io.BytesIO(docx_data))
        doc_text = "\n".join([p.text for p in doc.paragraphs])
        assert "KẾ HOẠCH BÀI DẠY (GIÁO ÁN)" in doc_text
        assert "THỂ TÍCH KHỐI LĂNG TRỤ ĐỨNG" in doc_text

        # Kiểm tra file PowerPoint trong ZIP
        pptx_data = zf.read(pptx_names[0])
        prs = Presentation(io.BytesIO(pptx_data))
        assert len(prs.slides) == len(sample_slide_deck_day19["slides"])

        # Kiểm tra file Hướng dẫn trong ZIP
        txt_data = zf.read(txt_names[0]).decode("utf-8")
        assert "TRỌN BỘ TÀI LIỆU DẠY HỌC CHUẨN CÔNG VĂN 5512/BGDĐT" in txt_data
        assert "LessonAI-5512" in txt_data


def test_task_19_2_export_dual_package_zip_validation_errors():
    """
    Task 19.2: Kiểm tra bắt lỗi đầu vào API export-dual-package:
    - Khi không truyền cả plan và slide_deck -> HTTP 400
    """
    response = client.post("/api/v1/lesson-plan/export-dual-package", json={})
    assert response.status_code == 400
    assert "Cần cung cấp ít nhất Kế hoạch bài dạy" in response.json()["detail"]


def test_task_19_2_frontend_ui_export_buttons_presence():
    """
    Task 19.2: Kiểm tra sự hiện diện và tích hợp của cụm nút bấm tải về trên giao diện:
    1. HeaderNav.jsx: Có dropdown 'Tải Về' chứa 'Tải Giáo Án Word (.docx)', 'Tải Bài Giảng PowerPoint (.pptx)', 'Tải Trọn Bộ Dual-Pack (.zip)'
    2. LessonPlanScreen.jsx: Có cụm nút 'Tải Word (.docx)', 'Tải Slide (.pptx)', 'Tải Trọn Bộ (.zip)'
    3. SlideEditorScreen.jsx: Có cụm nút 'Tải Slide (.pptx)', 'Tải Word (.docx)', 'Tải Trọn Bộ (.zip)'
    4. App.jsx: Có các handler handleExportDocx, handleExportPptx, handleExportDualPackage được truyền đầy đủ props
    """
    frontend_dir = ROOT_DIR / "app" / "frontend" / "src"

    # 1. HeaderNav.jsx
    header_path = frontend_dir / "components" / "HeaderNav.jsx"
    assert header_path.exists()
    header_content = header_path.read_text(encoding="utf-8")
    assert "Tải Về" in header_content
    assert "Tải Giáo Án Word (.docx)" in header_content
    assert "Tải Bài Giảng PowerPoint (.pptx)" in header_content
    assert "Tải Trọn Bộ Dual-Pack (.zip)" in header_content
    assert "onExportDualPackage" in header_content
    assert "onExportDocx" in header_content
    assert "onExportPptx" in header_content

    # 2. LessonPlanScreen.jsx
    lesson_path = frontend_dir / "components" / "LessonPlanScreen.jsx"
    assert lesson_path.exists()
    lesson_content = lesson_path.read_text(encoding="utf-8")
    assert "Tải Word (.docx)" in lesson_content
    assert "Tải Slide (.pptx)" in lesson_content
    assert "Tải Trọn Bộ (.zip)" in lesson_content
    assert "onExportDualPackage" in lesson_content
    assert "onExportPptx" in lesson_content

    # 3. SlideEditorScreen.jsx
    slide_editor_path = frontend_dir / "components" / "SlideEditorScreen.jsx"
    assert slide_editor_path.exists()
    slide_editor_content = slide_editor_path.read_text(encoding="utf-8")
    assert "Tải Slide (.pptx)" in slide_editor_content
    assert "Tải Word (.docx)" in slide_editor_content
    assert "Tải Trọn Bộ (.zip)" in slide_editor_content
    assert "onExportDualPackage" in slide_editor_content
    assert "onExportDocx" in slide_editor_content

    # 4. App.jsx
    app_path = frontend_dir / "App.jsx"
    assert app_path.exists()
    app_content = app_path.read_text(encoding="utf-8")
    assert "handleExportDualPackage" in app_content
    assert "handleExportDocx" in app_content
    assert "handleExportPptx" in app_content
    assert "/api/v1/lesson-plan/export-dual-package" in app_content
    assert "onExportDualPackage={handleExportDualPackage}" in app_content


# ==============================================================================
# TASK 19.3: TỐI ƯU LUỒNG DỮ LIỆU DUAL-MODE & HOÀN THIỆN ENDPOINT TẢI GÓI KÉP
# ==============================================================================

def test_task_19_3_dual_package_exporter_direct(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.3: Kiểm tra trực tiếp hàm export_dual_package:
    - Trả về đối tượng DualPackageResult với đầy đủ các trường dữ liệu
    - File Count = 3 (Word, PPTX, Manifest)
    - Checksum SHA-256 là chuỗi hex 64 ký tự hợp lệ
    - Manifest text chứa tiêu đề chuẩn và SHA-256 xác thực
    """
    payload = {
        "plan": sample_lesson_plan,
        "slide_deck": sample_slide_deck_day19
    }
    result = export_dual_package(payload)

    assert isinstance(result, DualPackageResult)
    assert len(result.zip_bytes) > 20000
    assert result.filename.endswith(".zip")
    assert result.file_count == 3
    assert len(result.docx_sha256) == 64
    assert len(result.pptx_sha256) == 64
    assert len(result.zip_sha256) == 64

    # Kiểm tra nội dung Manifest
    assert "TRỌN BỘ TÀI LIỆU DẠY HỌC CHUẨN CÔNG VĂN 5512/BGDĐT" in result.manifest_text
    assert result.docx_sha256 in result.manifest_text
    assert result.pptx_sha256 in result.manifest_text
    assert "OMML" in result.manifest_text
    assert "16:9" in result.manifest_text


def test_task_19_3_vietnamese_diacritics_sanitizer():
    """
    Task 19.3: Kiểm tra bộ tiền xử lý chuỗi tiếng Việt (remove_vietnamese_diacritics & sanitize_filename):
    - Khử dấu tiếng Việt chuẩn xác cho tất cả nguyên âm và ký tự Đ/đ
    - Chống lỗi mojibake / encoding khi giải nén trên hệ điều hành Windows cũ
    """
    raw_title = "Thể Tích Khối Lăng Trụ Đứng & Ứng Dụng Thực Tiễn (Toán 11)"
    cleaned = sanitize_filename(raw_title)
    assert "The_Tich_Khoi_Lang_Tru_Dung" in cleaned
    assert "Toan_11" in cleaned
    assert " " not in cleaned
    assert "&" not in cleaned
    assert "(" not in cleaned
    assert ")" not in cleaned

    # Trường hợp chuỗi rỗng
    assert sanitize_filename("") == "LessonAI_5512"
    assert sanitize_filename("   ") == "LessonAI_5512"


def test_task_19_3_concurrency_vs_sequential_execution(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.3: Kiểm tra tính toàn vẹn giữa chế độ xuất đa luồng (use_concurrency=True) và đơn luồng:
    - Cả hai chế độ đều sinh ra gói ZIP hợp lệ và cùng cấu trúc 3 tệp tin
    """
    payload = {
        "plan": sample_lesson_plan,
        "slide_deck": sample_slide_deck_day19
    }
    res_concurrent = export_dual_package(payload, use_concurrency=True)
    res_sequential = export_dual_package(payload, use_concurrency=False)

    assert res_concurrent.file_count == 3
    assert res_sequential.file_count == 3
    assert res_concurrent.docx_filename == res_sequential.docx_filename
    assert res_concurrent.pptx_filename == res_sequential.pptx_filename


def test_task_19_3_endpoint_custom_headers_and_integrity(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.3: Kiểm tra các custom headers phản hồi từ POST /api/v1/lesson-plan/export-dual-package:
    - X-Dual-Package-Docs-Count: '3'
    - X-Dual-Package-SHA256: khớp với mã hash của dữ liệu nhị phân trả về
    - X-Docx-SHA256 và X-Pptx-SHA256: độ dài 64 ký tự
    """
    payload = {
        "plan": sample_lesson_plan.model_dump(),
        "slide_deck": sample_slide_deck_day19
    }
    response = client.post("/api/v1/lesson-plan/export-dual-package", json=payload)
    assert response.status_code == 200
    headers = response.headers

    assert headers.get("x-dual-package-docs-count") == "3"
    assert "x-dual-package-sha256" in headers
    assert len(headers["x-dual-package-sha256"]) == 64
    assert len(headers["x-docx-sha256"]) == 64
    assert len(headers["x-pptx-sha256"]) == 64

    # Kiểm tra mã băm khớp với nội dung nhị phân
    import hashlib
    actual_hash = hashlib.sha256(response.content).hexdigest()
    assert headers["x-dual-package-sha256"] == actual_hash


def test_task_19_3_partial_payload_plan_only_or_deck_only(sample_lesson_plan, sample_slide_deck_day19):
    """
    Task 19.3: Kiểm tra khả năng xử lý linh hoạt khi người dùng chỉ cung cấp Kế hoạch bài dạy HOẶC Slide:
    - Chỉ có plan: Gói ZIP sinh ra 2 tệp (Word + Manifest)
    - Chỉ có slide_deck: Gói ZIP sinh ra 2 tệp (PPTX + Manifest)
    - Hệ thống không bị crash và trả về HTTP 200
    """
    # 1. Chỉ có plan
    resp_plan = client.post("/api/v1/lesson-plan/export-dual-package", json={"plan": sample_lesson_plan.model_dump()})
    assert resp_plan.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp_plan.content), "r") as zf:
        names = zf.namelist()
        assert len(names) == 2
        assert any(n.endswith(".docx") for n in names)
        assert any(n.endswith(".txt") for n in names)
        assert not any(n.endswith(".pptx") for n in names)

    # 2. Chỉ có slide_deck
    resp_deck = client.post("/api/v1/lesson-plan/export-dual-package", json={"slide_deck": sample_slide_deck_day19})
    assert resp_deck.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp_deck.content), "r") as zf:
        names = zf.namelist()
        assert len(names) == 2
        assert any(n.endswith(".pptx") for n in names)
        assert any(n.endswith(".txt") for n in names)
        assert not any(n.endswith(".docx") for n in names)


# ==============================================================================
# TASK 19.4: SCRIPT DỌN DẸP FILE TẠM TRÊN SERVER & QUẢN LÝ BỘ NHỚ
# ==============================================================================

def test_task_19_4_cleaner_detects_and_deletes_expired_files(tmp_path):
    """
    Task 19.4: Kiểm tra khả năng phát hiện và xóa chính xác các tệp tạm hết hạn:
    - Các tệp tạm hết hạn (> 3600s) bị xóa
    - Các tệp tạm mới sinh (< 3600s) được bảo lưu an toàn
    - Các tệp mã nguồn (.py) và cấu hình (.json) KHÔNG BAO GIỜ bị xóa nhầm
    """
    now = time.time()

    # 1. Tệp tạm hết hạn (> 1 giờ)
    expired_chart = tmp_path / "chart_parabol_temp.png"
    expired_chart.write_bytes(b"FAKE_PNG_BYTES_FOR_TEST_123456789")
    os.utime(expired_chart, (now - 7200, now - 7200))

    expired_docx = tmp_path / "temp_giao_an_expired.docx"
    expired_docx.write_bytes(b"FAKE_DOCX_BYTES_FOR_TEST_987654321")
    os.utime(expired_docx, (now - 5000, now - 5000))

    # 2. Tệp tạm mới sinh (< 1 giờ)
    fresh_tmp = tmp_path / "temp_slide_current.tmp"
    fresh_tmp.write_bytes(b"FRESH_TEMP_BYTES")
    os.utime(fresh_tmp, (now - 60, now - 60))

    # 3. Tệp mã nguồn và cấu hình quan trọng (Bảo vệ tuyệt đối)
    source_file = tmp_path / "important_service.py"
    source_file.write_text("print('HELLO WORLD DO NOT DELETE')", encoding="utf-8")
    os.utime(source_file, (now - 10000, now - 10000))

    config_file = tmp_path / "system_config.json"
    config_file.write_text('{"setting": "keep"}', encoding="utf-8")
    os.utime(config_file, (now - 10000, now - 10000))

    cleaner = TempFileCleaner(target_dirs=[str(tmp_path)])
    report = cleaner.clean(max_age_seconds=3600, force_all=False, dry_run=False, clean_memory=False)

    assert isinstance(report, CleaningReport)
    assert report.deleted_files == 2
    assert report.bytes_reclaimed > 0

    # Xác thực trên ổ đĩa
    assert not expired_chart.exists(), "Tệp đồ thị tạm hết hạn phải bị xóa"
    assert not expired_docx.exists(), "Tệp Word tạm hết hạn phải bị xóa"
    assert fresh_tmp.exists(), "Tệp tạm mới sinh phải được giữ lại"
    assert source_file.exists(), "Tệp mã nguồn .py TUYỆT ĐỐI không bị xóa"
    assert config_file.exists(), "Tệp cấu hình .json TUYỆT ĐỐI không bị xóa"


def test_task_19_4_force_all_and_dry_run_modes(tmp_path):
    """
    Task 19.4: Kiểm tra chế độ force_all và dry_run:
    - dry_run=True: Chỉ mô phỏng tính toán, không xóa thật trên đĩa
    - force_all=True: Xóa toàn bộ tệp tạm không kể tuổi tệp
    """
    file1 = tmp_path / "chart_1.png"
    file1.write_bytes(b"12345")
    file2 = tmp_path / "temp_notes.tmp"
    file2.write_bytes(b"67890")

    cleaner = TempFileCleaner(target_dirs=[str(tmp_path)])

    # 1. Chạy thử dry_run
    dry_report = cleaner.clean(max_age_seconds=3600, force_all=True, dry_run=True, clean_memory=False)
    assert dry_report.deleted_files == 2
    assert file1.exists(), "dry_run không được xóa tệp thật"
    assert file2.exists(), "dry_run không được xóa tệp thật"

    # 2. Xóa thật với force_all=True
    real_report = cleaner.clean(max_age_seconds=3600, force_all=True, dry_run=False, clean_memory=False)
    assert real_report.deleted_files == 2
    assert not file1.exists()
    assert not file2.exists()


def test_task_19_4_temp_stats_reporting(tmp_path):
    """
    Task 19.4: Kiểm tra tra cứu thống kê số lượng và dung lượng tệp tin tạm (scan_stats)
    """
    cleaner = TempFileCleaner(target_dirs=[str(tmp_path)])

    # Ban đầu thư mục trống
    stats_empty = cleaner.scan_stats()
    assert stats_empty.total_files == 0
    assert stats_empty.total_bytes == 0

    # Tạo 2 tệp tạm với dung lượng biết trước
    f1 = tmp_path / "temp_file_a.png"
    f1.write_bytes(b"A" * 1024)  # 1 KB

    f2 = tmp_path / "temp_file_b.docx"
    f2.write_bytes(b"B" * 2048)  # 2 KB

    stats = cleaner.scan_stats()
    assert stats.total_files == 2
    assert stats.total_bytes == 3072
    assert stats.total_mb > 0
    assert stats.oldest_file_age_seconds is not None


def test_task_19_4_memory_and_matplotlib_cleanup():
    """
    Task 19.4: Kiểm tra giải phóng tài nguyên bộ nhớ rò rỉ:
    - Mở các figure Matplotlib rồi gọi hàm dọn dẹp
    - Đảm bảo tất cả figure được đóng hoàn toàn
    - Kích hoạt Python Garbage Collection thu gom rác
    """
    import matplotlib.pyplot as plt

    # Tạo 3 figure tạm để kiểm tra
    fig1 = plt.figure()
    fig2 = plt.figure()
    fig3 = plt.figure()

    assert len(plt.get_fignums()) >= 3

    # Kích hoạt dọn dẹp bộ nhớ
    res = cleanup_memory_and_figures()

    assert len(plt.get_fignums()) == 0, "Tất cả Matplotlib figures phải được đóng triệt để"
    assert res["figures_closed"] >= 3
    assert "collected_objects" in res
    assert "current_memory_mb" in res


def test_task_19_4_system_api_endpoints():
    """
    Task 19.4: Kiểm tra các API quản trị và bảo trì hệ thống:
    - POST /api/v1/system/clean-temp-files
    - GET /api/v1/system/temp-stats
    - POST /api/v1/system/cleanup-memory
    """
    # 1. Thống kê tệp tạm
    resp_stats = client.get("/api/v1/system/temp-stats")
    assert resp_stats.status_code == 200
    data_stats = resp_stats.json()
    assert "total_files" in data_stats
    assert "total_bytes" in data_stats
    assert "directories" in data_stats

    # 2. Dọn dẹp tệp tạm (chế độ dry_run)
    resp_clean = client.post("/api/v1/system/clean-temp-files?dry_run=true&max_age_seconds=1800")
    assert resp_clean.status_code == 200
    data_clean = resp_clean.json()
    assert "scanned_files" in data_clean
    assert "deleted_files" in data_clean
    assert "bytes_reclaimed" in data_clean
    assert "execution_time_ms" in data_clean

    # 3. Giải phóng bộ nhớ
    resp_mem = client.post("/api/v1/system/cleanup-memory")
    assert resp_mem.status_code == 200
    data_mem = resp_mem.json()
    assert "figures_closed" in data_mem
    assert "collected_objects" in data_mem
    assert "current_memory_mb" in data_mem


# ==============================================================================
# TASK 19.5: NGHIỆM THU DEFINITION OF DONE (DoD) TOÀN DIỆN NGÀY 19
# ==============================================================================

def test_task_19_5_definition_of_done_comprehensive_audit(sample_lesson_plan, sample_slide_deck_day19, tmp_path):
    """
    Task 19.5: Nghiệm thu toàn diện Definition of Done (DoD) Ngày 19:
    1. Chuẩn hóa Word (.docx): Khổ A4, lề trang hành chính NĐ 30/2020, font Times New Roman,
       công thức OMML native XML, lọc sạch ký tự điều khiển cấm.
    2. Đồng bộ 100% Dual-Mode: Hoạt động 3 nhúng câu hỏi trắc nghiệm & bài giải mẫu tự luận;
       Mục II bổ sung thiết bị 3D WebGL; Hoạt động 2 nhúng công thức định nghĩa.
    3. Gói kép Dual-Package ZIP: Tệp tin nén tối ưu, khử dấu tiếng Việt tên file,
       đính kèm file Hướng dẫn sư phạm (Manifest) và mã băm SHA-256 xác thực.
    4. Quản lý tài nguyên & Dọn dẹp tệp tạm: Xóa tệp tạm hết hạn, bảo vệ code nguồn,
       giải phóng Matplotlib figures và bộ nhớ rò rỉ.
    5. Giao diện Web Studio: Cụm nút bấm xuất bản hiện diện trên cả 3 màn hình và Header.
    """
    # --------------------------------------------------------------------------
    # MỤC 1: KIỂM ĐỊNH TỆP WORD (.DOCX) CHUẨN CÔNG VĂN 5512 & NGHỊ ĐỊNH 30/2020
    # --------------------------------------------------------------------------
    docx_stream = export_lesson_plan_to_docx(sample_lesson_plan, slide_deck=sample_slide_deck_day19)
    assert isinstance(docx_stream, io.BytesIO)
    assert docx_stream.getbuffer().nbytes > 15000

    doc = Document(docx_stream)

    # 1.1 Kiểm tra kích thước trang A4 & lề trang hành chính
    section = doc.sections[0]
    assert section.page_width is not None and section.page_height is not None
    assert section.left_margin is not None and section.right_margin is not None
    assert section.top_margin is not None and section.bottom_margin is not None
    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69
    assert round(section.left_margin.inches, 2) == 1.18     # 3.0 cm
    assert round(section.right_margin.inches, 2) == 0.59    # 1.5 cm
    assert round(section.top_margin.inches, 2) == 0.79      # 2.0 cm
    assert round(section.bottom_margin.inches, 2) == 0.79   # 2.0 cm

    # 1.2 Kiểm tra sự hiện diện của công thức toán OMML trong XML gốc của Word
    doc_xml_str = doc._element.xml
    assert "oMath" in doc_xml_str or "m:oMath" in doc_xml_str or "Cambria Math" in doc_xml_str, \
        "Tài liệu Word phải chứa thẻ công thức toán học OMML hoặc font toán học Cambria Math"

    # 1.3 Kiểm tra lọc sạch ký tự điều khiển cấm (không chứa ASCII 0-8, 11-12, 14-31)
    for p in doc.paragraphs:
        for ch in p.text:
            code = ord(ch)
            assert not (0 <= code <= 8 or 11 <= code <= 12 or 14 <= code <= 31), \
                f"Phát hiện ký tự cấm ASCII {code} trong văn bản Word"

    # --------------------------------------------------------------------------
    # MỤC 2: KIỂM ĐỊNH TÍNH ĐỒNG BỘ 100% NỘI DUNG SƯ PHẠM (DUAL-MODE SYNC)
    # --------------------------------------------------------------------------
    full_text = "\n".join([p.text for p in doc.paragraphs])

    # 2.1 Câu hỏi trắc nghiệm đồng bộ ở Hoạt động 3
    assert "HỆ THỐNG CÂU HỎI TRẮC NGHIỆM ĐÁNH GIÁ ĐỊNH DẠNG BGD&ĐT" in full_text
    assert "Cho hình lăng trụ đứng có diện tích đáy" in full_text
    assert "[A]." in full_text and "[B]." in full_text and "[C]." in full_text and "[D]." in full_text
    assert "[Đáp án đúng chuẩn SGK]: Phương án B" in full_text
    assert "Nhầm công thức thể tích khối chóp có hệ số 1/3" in full_text

    # 2.2 Lời giải tự luận mẫu 3 bước
    assert "HƯỚNG DẪN GIẢI MẪU TỰ LUẬN TỪNG BƯỚC CHUẨN MỰC SGK:" in full_text
    assert "Bước 1: Tính diện tích đáy tam giác đều" in full_text
    assert "Bước 2: Xác định chiều cao lăng trụ đứng" in full_text
    assert "Bước 3: Thể tích khối lăng trụ" in full_text

    # 2.3 Thiết bị dạy học bổ sung 3D WebGL và Studio ở Mục II
    assert "Mô hình Hình học không gian 3D tương tác WebGL" in full_text
    assert "Bài giảng điện tử tương tác chuẩn Web Studio" in full_text

    # --------------------------------------------------------------------------
    # MỤC 3: KIỂM ĐỊNH GÓI KÉP DUAL-PACKAGE (.ZIP) QUA API
    # --------------------------------------------------------------------------
    dual_payload = {
        "plan": sample_lesson_plan.model_dump(),
        "slide_deck": sample_slide_deck_day19
    }
    dual_resp = client.post("/api/v1/lesson-plan/export-dual-package", json=dual_payload)
    assert dual_resp.status_code == 200
    assert "application/zip" in dual_resp.headers["content-type"]
    assert dual_resp.headers.get("x-dual-package-docs-count") == "3"

    zip_bytes = io.BytesIO(dual_resp.content)
    with zipfile.ZipFile(zip_bytes, "r") as zf:
        namelist = zf.namelist()
        assert len(namelist) == 3

        # Kiểm tra tên tệp tin không chứa dấu tiếng Việt (chống lỗi mojibake)
        for fname in namelist:
            assert all(ord(c) < 128 for c in fname), f"Tên tệp {fname} trong ZIP phải là ASCII an toàn"

        # Kiểm tra file Hướng dẫn Sư phạm (Manifest)
        manifest_data = zf.read("Huong_Dan_Goi_Kep_5512.txt").decode("utf-8")
        assert "TRỌN BỘ TÀI LIỆU DẠY HỌC CHUẨN CÔNG VĂN 5512/BGDĐT" in manifest_data
        assert "KẾ HOẠCH BÀI DẠY (GIÁO ÁN WORD):" in manifest_data
        assert "BÀI GIẢNG ĐIỆN TỬ (POWERPOINT PPTX):" in manifest_data
        assert "Checksum (SHA-256):" in manifest_data

    # --------------------------------------------------------------------------
    # MỤC 4: KIỂM ĐỊNH BẢO VỆ TÀI NGUYÊN & DỌN DẸP BỘ NHỚ
    # --------------------------------------------------------------------------
    # Đảm bảo sau khi xuất gói kép, Matplotlib không bị rò rỉ figure
    import matplotlib.pyplot as plt
    assert len(plt.get_fignums()) == 0, "Không được để sót Matplotlib figure sau khi gọi API"

    # Kiểm tra hệ thống dọn dẹp tệp tạm hoạt động ổn định
    cleaner_report = clean_expired_temp_files(max_age_seconds=7200, dry_run=True)
    assert cleaner_report.errors == []



