from io import BytesIO
from pathlib import Path
from typing import List, Tuple, Any, Dict, Union, BinaryIO, TypeAlias
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Định nghĩa bí danh kiểu dữ liệu đầu vào cho file PowerPoint (PEP 613 TypeAlias)
Source: TypeAlias = Union[str, Path, bytes, BinaryIO, BytesIO]

try:
    from .common import read_source, clean_text
except ImportError:
    from parsers.common import read_source, clean_text


def _extract_shapes_text(
    shapes: Any,
    skip_shape: Any = None
) -> Tuple[List[str], List[List[List[str]]]]:
    """
    Duyệt đệ quy qua danh sách shapes của một slide để bóc tách:
    1. paragraphs (List[str]): Danh sách các đoạn văn bản đã làm sạch.
    2. tables (List[List[List[str]]]): Danh sách các bảng (mỗi bảng gồm các hàng -> các ô).
    """
    paragraphs: List[str] = []
    tables: List[List[List[str]]] = []

    for shape in shapes:
        # Bỏ qua shape tiêu đề nếu đã được lấy riêng ở slide.shapes.title
        if skip_shape is not None and shape == skip_shape:
            continue

        # 1. TRƯỜNG HỢP ĐỆ QUY: Nếu là GroupShape -> duyệt sâu vào các shape con bên trong
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            sub_paragraphs, sub_tables = _extract_shapes_text(
                shape.shapes, skip_shape=skip_shape
            )
            paragraphs.extend(sub_paragraphs)
            tables.extend(sub_tables)
            continue

        # 2. TRƯỜNG HỢP KHUNG VĂN BẢN (Text Frame)
        tf = getattr(shape, "text_frame", None)
        if tf is not None:
            for para in getattr(tf, "paragraphs", []):
                cleaned_para = clean_text(getattr(para, "text", ""))
                if cleaned_para:
                    paragraphs.append(cleaned_para)

        # 3. TRƯỜNG HỢP BẢNG BIỂU (Table)
        if getattr(shape, "has_table", False):
            table_rows = _extract_table(shape.table)
            if table_rows:
                tables.append(table_rows)

    return paragraphs, tables

def _extract_table(table: Any) -> List[List[str]]:
    """
    Chuyển đối tượng bảng (Table) của python-pptx thành ma trận 2 chiều list[list[str]]:
    - Làm sạch văn bản trong từng ô bằng clean_text().
    - Xử lý ô bị gộp (is_spanned) để không lặp nội dung.
    - Loại bỏ các hàng hoàn toàn trống.
    """
    table_rows: List[List[str]] = []

    for row in table.rows:
        row_cells: List[str] = []
        for cell in row.cells:
            # Nếu ô này bị gộp chìm vào ô khác (is_spanned = True) -> để chuỗi rỗng tránh lặp
            if getattr(cell, "is_spanned", False):
                row_cells.append("")
            else:
                row_cells.append(clean_text(cell.text))

        # Chỉ giữ lại hàng nếu có ít nhất 1 ô có nội dung khác rỗng
        if any(row_cells):
            table_rows.append(row_cells)

    return table_rows

def _get_title(slide: Any) -> str:
    """
    Lấy tiêu đề của một trang Slide PowerPoint:
    - Ưu tiên 1: Lấy từ placeholder tiêu đề chuẩn (slide.shapes.title).
    - Ưu tiên 2 (Fallback): Nếu slide không có title placeholder (hoặc bị bỏ trống),
      tìm shape văn bản nằm ở vị trí cao nhất trên slide (shape.top nhỏ nhất)
      và lấy dòng đầu tiên làm tiêu đề.
    """
    # 1. Kiểm tra Title Placeholder chuẩn của PowerPoint
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        title_text = clean_text(getattr(title_shape, "text", ""))
        if title_text:
            return title_text

    # 2. Fallback: Tìm trong các text shape có sẵn (sắp xếp theo tọa độ top từ trên xuống dưới)
    text_shapes = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and clean_text(getattr(shape, "text", ""))
    ]
    if text_shapes:
        # Sắp xếp theo vị trí từ cao xuống thấp (top nhỏ nhất là nằm trên cùng của slide)
        text_shapes.sort(key=lambda s: (getattr(s, "top", 0) or 0))
        first_text = clean_text(getattr(text_shapes[0], "text", ""))
        # Chỉ lấy dòng đầu tiên nếu hộp văn bản đó có nhiều dòng
        return first_text.split("\n")[0]

    return ""

def _get_notes(slide: Any) -> str:
    """
    Trích xuất lời thoại / ghi chú của giáo viên (Speaker Notes) trong một trang Slide:
    - Kiểm tra slide.has_notes_slide trước để tránh python-pptx tự tạo XML rỗng.
    - Kiểm tra an toàn notes_text_frame và làm sạch văn bản bằng clean_text().
    """
    if not getattr(slide, "has_notes_slide", False):
        return ""

    try:
        notes_slide = getattr(slide, "notes_slide", None)
        if notes_slide is not None:
            notes_tf = getattr(notes_slide, "notes_text_frame", None)
            if notes_tf is not None:
                return clean_text(getattr(notes_tf, "text", ""))
    except Exception:
        # Phòng vệ trường hợp cấu trúc XML của notes_slide trong file cũ bị lỗi hỏng
        return ""

    return ""


def parse_pptx(source: Source) -> Dict[str, Any]:
    """
    Hàm chính bóc tách toàn bộ nội dung file PowerPoint (.pptx):
    - Đầu vào: Đường dẫn file (str | Path), bytes, hoặc BytesIO/BinaryIO.
    - Đầu ra: Dictionary chứa tiêu đề bài giảng, tổng số slide, danh sách chi tiết từng slide
      và văn bản tổng hợp (raw_text) đã làm sạch.
    """
    stream = read_source(source)
    prs = Presentation(stream)

    slides_data: List[Dict[str, Any]] = []
    raw_blocks: List[str] = []
    presentation_title: str = ""

    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = getattr(slide.shapes, "title", None)
        title = _get_title(slide)

        paragraphs, tables = _extract_shapes_text(
            slide.shapes, skip_shape=title_shape
        )

        # Nếu tiêu đề được lấy từ nhánh Fallback (trùng với đoạn văn đầu tiên),
        # loại bỏ đoạn trùng để không bị lặp lại tiêu đề trong phần thân
        if title and paragraphs and paragraphs[0] == title:
            paragraphs = paragraphs[1:]

        notes = _get_notes(slide)

        # Lưu tiêu đề của slide đầu tiên có chữ làm tiêu đề chung của bài giảng
        if not presentation_title and title:
            presentation_title = title

        # Đóng gói dữ liệu cấu trúc của Slide
        slides_data.append({
            "slide_index": idx,
            "title": title,
            "paragraphs": paragraphs,
            "tables": tables,
            "notes": notes,
        })

        # Ghép văn bản tổng hợp (raw_text) cho AI
        block_lines: List[str] = [
            f"--- SLIDE {idx}{f': {title}' if title else ''} ---"
        ]
        if paragraphs:
            block_lines.extend(paragraphs)
        for table in tables:
            for row in table:
                block_lines.append(" | ".join(row))
        if notes:
            block_lines.append(f"[Ghi chú GV]: {notes}")

        raw_blocks.append("\n".join(block_lines))

    # Fallback tiêu đề từ metadata của file nếu tất cả các slide đều không có title
    if not presentation_title:
        meta_title = clean_text(getattr(prs.core_properties, "title", ""))
        presentation_title = meta_title or "Bài giảng PowerPoint"

    return {
        "file_type": "pptx",
        "title": presentation_title,
        "total_slides": len(slides_data),
        "slides": slides_data,
        "raw_text": "\n\n".join(raw_blocks),
    }
