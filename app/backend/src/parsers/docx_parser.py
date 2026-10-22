from io import BytesIO
from pathlib import Path
from typing import List, Dict, Any, Iterator, Union, BinaryIO, TypeAlias
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

# Định nghĩa bí danh kiểu dữ liệu đầu vào cho file Word (PEP 613 TypeAlias)
Source: TypeAlias = Union[str, Path, bytes, BinaryIO, BytesIO]

try:
    from .common import read_source, clean_text
except ImportError:
    from parsers.common import read_source, clean_text


def _dedupe_row_cells(row: Any) -> List[str]:
    """
    Trích xuất danh sách văn bản các ô (cells) trên một hàng của bảng Word,
    đồng thời khử trùng lặp các ô bị gộp ngang/dọc (Merged Cells):
    - Kiểm tra định danh thẻ XML gốc `cell._tc`.
    - Nếu thẻ `_tc` đã xuất hiện trong hàng -> bỏ qua để không bị nhân đôi văn bản.
    - Làm sạch văn bản từng ô bằng clean_text().
    """
    cells_text: List[str] = []
    seen_tc = set()

    for cell in row.cells:
        tc = getattr(cell, "_tc", None)
        if tc is not None:
            if tc in seen_tc:
                continue
            seen_tc.add(tc)

        cells_text.append(clean_text(cell.text))

    return cells_text

def _table_to_rows(table: Any) -> List[List[str]]:
    """
    Chuyển một bảng (Table) trong Word thành danh sách các hàng list[list[str]]:
    - Sử dụng _dedupe_row_cells() để lấy nội dung từng hàng đã khử ô gộp ngang.
    - Bỏ qua các hàng hoàn toàn trống.
    - Khử trùng lặp các hàng liên tiếp giống hệt nhau (do gộp ô dọc trong Word).
    """
    rows: List[List[str]] = []

    for row in table.rows:
        row_cells = _dedupe_row_cells(row)

        # 1. Bỏ qua hàng nếu tất cả các ô trong hàng đều trống
        if not any(row_cells):
            continue

        # 2. Bỏ qua nếu hàng này lặp lại y hệt hàng ngay trước đó (do vertical merge)
        if rows and row_cells == rows[-1]:
            continue

        rows.append(row_cells)

    return rows

def _detect_title(doc: Any) -> str:
    """
    Tự động nhận diện tiêu đề bài học từ đối tượng Document của python-docx:
    - Ưu tiên 1: Đoạn văn có style 'Title'/'Heading 1' hoặc bắt đầu bằng từ khóa bài học
      ('Tên bài dạy', 'Bài ', 'Chủ đề', 'Chuyên đề').
    - Ưu tiên 2: Đoạn văn có nội dung đầu tiên trong văn bản.
    - Ưu tiên 3 (Fallback): Ô đầu tiên trong bảng hoặc metadata của file Word.
    """
    non_empty_paras: List[str] = []
    lesson_keywords = ("tên bài dạy", "bài dạy:", "bài ", "chủ đề", "chuyên đề")

    for para in doc.paragraphs:
        text = clean_text(para.text)
        if not text:
            continue

        non_empty_paras.append(text)
        style_name = getattr(getattr(para, "style", None), "name", "") or ""
        lower_text = text.lower()

        # Ưu tiên 1a: Style là Title hoặc Heading 1
        if style_name.lower() in ("title", "heading 1"):
            return text

        # Ưu tiên 1b: Bắt đầu bằng từ khóa tên bài dạy chuẩn CV 5512
        if lower_text.startswith(lesson_keywords):
            return text

    # Ưu tiên 2: Đoạn văn không rỗng đầu tiên
    if non_empty_paras:
        return non_empty_paras[0]

    # Ưu tiên 3 (Fallback): Lấy từ bảng đầu tiên nếu văn bản nằm hoàn toàn trong bảng
    for table in doc.tables:
        rows = _table_to_rows(table)
        if rows and rows[0]:
            for cell_text in rows[0]:
                if cell_text:
                    return cell_text.split("\n")[0]

    # Fallback cuối cùng: Metadata của file Word
    meta_title = clean_text(getattr(doc.core_properties, "title", ""))
    return meta_title or "Giáo án Word"

def _iter_block_items(parent: Any) -> Iterator[Union[Paragraph, Table]]:
    """
    Bộ sinh (Generator) duyệt tuần tự các khối nội dung (Paragraph và Table)
    theo đúng thứ tự xuất hiện từ trên xuống dưới trong tài liệu Word:
    - Nếu thẻ XML con là CT_P (<w:p>) -> trả về đối tượng Paragraph.
    - Nếu thẻ XML con là CT_Tbl (<w:tbl>) -> trả về đối tượng Table.
    """
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        parent_elm = parent.element.body

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def parse_docx(source: Source) -> Dict[str, Any]:
    """
    Hàm chính bóc tách toàn bộ nội dung file Giáo án Word (.docx):
    - Đầu vào: Đường dẫn file (str | Path), bytes, hoặc BytesIO/BinaryIO.
    - Đầu ra: Dictionary chứa tiêu đề bài học, danh sách đoạn văn, danh sách bảng biểu,
      danh sách khối nội dung tuần tự (blocks) và văn bản tổng hợp (raw_text) đã làm sạch.
    """
    stream = read_source(source)
    doc = Document(stream)

    title = _detect_title(doc)

    paragraphs: List[str] = []
    tables: List[List[List[str]]] = []
    blocks: List[Dict[str, Any]] = []
    raw_lines: List[str] = []

    for item in _iter_block_items(doc):
        if isinstance(item, Paragraph):
            text = clean_text(item.text)
            if text:
                paragraphs.append(text)
                blocks.append({"type": "paragraph", "text": text})
                raw_lines.append(text)

        elif isinstance(item, Table):
            rows = _table_to_rows(item)
            if rows:
                tables.append(rows)
                blocks.append({"type": "table", "rows": rows})
                table_text = "\n".join(" | ".join(row) for row in rows)
                raw_lines.append(f"[BẢNG]:\n{table_text}")

    return {
        "file_type": "docx",
        "title": title,
        "total_paragraphs": len(paragraphs),
        "total_tables": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "blocks": blocks,
        "raw_text": "\n\n".join(raw_lines),
    }
